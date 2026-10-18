"""读取同一份内存快照，按已确认映射转换为 RawCsvRow，复用既有校验器。"""

import hashlib
import io
import posixpath
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.styles.numbers import is_datetime
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from openpyxl.utils.datetime import CALENDAR_WINDOWS_1900

from csv_reader import CsvValidationError, RawCsvRow
from field_mapping import MappingError, MappingPlan, analyze_headers, resolve_mapping
from order_validator import OrderLine, validate_rows


class ExcelImportError(ValueError):
    """不返回部分合格数据；消息只含位置和原因，不输出完整订单。"""


@dataclass(frozen=True)
class CellValue:
    value: object
    kind: str
    number_format: str
    number_text: str | None = None
    epoch_1900: bool = False


@dataclass(frozen=True)
class SheetSnapshot:
    name: str
    rows: tuple[tuple[CellValue, ...], ...]
    merged: tuple[str, ...]
    hidden_rows: tuple[int, ...]
    hidden_columns: tuple[tuple[int, int], ...]
    state: str


@dataclass(frozen=True)
class ExcelSnapshot:
    sha256: str
    sheets: tuple[SheetSnapshot, ...]


@dataclass(frozen=True)
class SheetSelection:
    snapshot_sha256: str
    sheet: SheetSnapshot
    header_row: int
    headers: tuple[str, ...]

    @property
    def context(self):
        return repr((self.snapshot_sha256, self.sheet.name, self.header_row))


def _problem(sheet, row, column, reason):
    location = f"工作表 {sheet!r}"
    if row is not None:
        location += f" 第 {row} 行"
    if column is not None:
        location += f" {get_column_letter(column)} 列"
    return ExcelImportError(f"{location}：{reason}")


def _check_raw_iso_date(text):
    """在库解析前验证 t=d 原文全文，不能依赖已经丢失时区或尾缀的 datetime。"""
    date_part = r"[0-9]{4}-[0-9]{2}-[0-9]{2}"
    time_part = r"[0-9]{2}:[0-9]{2}:[0-9]{2}(?:\.0+)?"
    if not re.fullmatch(rf"(?:{date_part}(?:T{time_part})?|{time_part})", text):
        raise ValueError("原生 ISO 日期不支持时区、非零小数秒或非法尾随内容，必须采用明确的无时区格式。")
    # 日期和时间单独出现是合法 XML 类型，但映射为 order_time 时仍由原有转换规则拒绝。
    try:
        if "T" in text:
            datetime.fromisoformat(text)
        elif ":" in text:
            time.fromisoformat(text)
        else:
            date.fromisoformat(text)
    except ValueError as error:
        raise ValueError("原生 ISO 日期或时间不存在，不能可靠解析。") from error


def _numeric_sources(archive, content_types):
    """保留数值字面量，并提前严格检查 ISO 日期原文；其他读取交给 openpyxl。"""
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    relation = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    manifest = ElementTree.fromstring(content_types)
    part = next(item.attrib["PartName"].lstrip("/") for item in manifest
                if item.attrib.get("ContentType") == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml")
    folder, filename = posixpath.split(part)
    relations = ElementTree.fromstring(archive.read(posixpath.join(folder, "_rels", filename + ".rels")))
    targets = {item.attrib["Id"]: item for item in relations}
    workbook_xml = ElementTree.fromstring(archive.read(part))
    result = {}
    for sheet in workbook_xml.findall(f"{{{main}}}sheets/{{{main}}}sheet"):
        target = targets[sheet.attrib[f"{{{relation}}}id"]]
        if not target.attrib.get("Type", "").endswith("/worksheet"):
            continue
        if target.attrib.get("TargetMode") == "External":
            raise ExcelImportError("不支持位于外部文件的工作表。")
        path = target.attrib["Target"]
        path = path.lstrip("/") if path.startswith("/") else posixpath.normpath(posixpath.join(folder, path))
        root = ElementTree.fromstring(archive.read(path))
        numbers = {}
        for cell in root.findall(f"{{{main}}}sheetData/{{{main}}}row/{{{main}}}c"):
            value = cell.find(f"{{{main}}}v")
            if cell.attrib.get("t") == "d":
                try:
                    _check_raw_iso_date(value.text if value is not None and value.text is not None else "")
                except ValueError as error:
                    row, column = coordinate_to_tuple(cell.attrib["r"])
                    raise _problem(sheet.attrib["name"], row, column, str(error)) from error
            if cell.attrib.get("t", "n") == "n" and value is not None and value.text is not None:
                numbers[cell.attrib["r"]] = value.text
        result[sheet.attrib["name"]] = numbers
    return result


def read_snapshot(path: str | Path) -> ExcelSnapshot:
    """唯一一次读取路径；此后不重新打开源文件，也不创建任何文件。"""
    workbook = None
    try:
        path = Path(path)
        if path.suffix.lower() != ".xlsx":
            raise ExcelImportError("只支持普通 .xlsx 文件，不支持 .xls、.xlsm 或更改后缀冒充的文件。")
        content = path.read_bytes()
        if content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            raise ExcelImportError("文件可能已加密或属于旧版 Excel 格式，当前不支持。")
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if any(item.flag_bits & 1 for item in archive.infolist()):
                raise ExcelImportError("不支持加密工作簿。")
            types = archive.read("[Content_Types].xml")
            if b"macroEnabled" in types or b"vbaProject" in types or b"spreadsheetml.template" in types:
                raise ExcelImportError("不支持含宏或模板的工作簿，仅支持普通 .xlsx。")
            numbers = _numeric_sources(archive, types)
        # False 保留公式本体，绝不能用上次 Excel 保存的缓存值作为交易数据。
        workbook = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
        sheets = []
        for sheet in workbook.worksheets:
            rows = tuple(tuple(CellValue(cell.value, cell.data_type, cell.number_format,
                                         numbers.get(sheet.title, {}).get(cell.coordinate),
                                         workbook.epoch == CALENDAR_WINDOWS_1900) for cell in row)
                         for row in sheet.iter_rows())
            sheets.append(SheetSnapshot(
                sheet.title, rows, tuple(str(area) for area in sheet.merged_cells.ranges),
                tuple(index for index, dimension in sheet.row_dimensions.items() if dimension.hidden),
                tuple((dimension.min, dimension.max) for dimension in sheet.column_dimensions.values()
                      if dimension.hidden), sheet.sheet_state,
            ))
        if not sheets:
            raise ExcelImportError("工作簿没有可读取的普通工作表。")
        return ExcelSnapshot(hashlib.sha256(content).hexdigest(), tuple(sheets))
    except ExcelImportError:
        raise
    except Exception as error:
        # 解析器错误可能包含单元格原文；只展示异常类型，避免泄露订单或客户资料。
        raise ExcelImportError(f"无法读取工作簿（{type(error).__name__}）：可能损坏、加密、权限不足或格式不支持。") from error
    finally:
        if workbook is not None:
            workbook.close()


def header_candidates(sheet: SheetSnapshot) -> tuple[tuple[int, int], ...]:
    """仅提示哪些行包含已知表头，不自动选择，不展示数据行的内容。"""
    result = []
    for index, row in enumerate(sheet.rows, 1):
        score = 0
        for cell in row:
            if cell.kind != "f" and isinstance(cell.value, str) and cell.value.strip():
                score += bool(analyze_headers((cell.value,))[0].candidates)
        if score:
            result.append((index, score))
    return tuple(result)


def select_sheet(snapshot: ExcelSnapshot, name: str, header_row: int) -> SheetSelection:
    sheet = next((item for item in snapshot.sheets if item.name == name), None)
    if sheet is None:
        raise _problem(name, None, None, "工作表不存在，必须明确选择。")
    if type(header_row) is not int or not 1 <= header_row <= len(sheet.rows):
        raise _problem(name, None, None, "表头行号无效。")
    if sheet.state != "visible":
        raise _problem(name, header_row, None, "所选工作表被隐藏，当前不支持。")
    if sheet.merged:
        raise _problem(name, header_row, None, "存在合并单元格 " + "、".join(sheet.merged) + "，不能可靠解释。")
    if any(row >= header_row for row in sheet.hidden_rows):
        row = next(row for row in sorted(sheet.hidden_rows) if row >= header_row)
        raise _problem(name, row, None, "存在隐藏行，不能静默排除或推断其用途。")
    if sheet.hidden_columns:
        raise _problem(name, header_row, sheet.hidden_columns[0][0], "存在隐藏列，不能静默忽略。")
    headers = []
    for column, cell in enumerate(sheet.rows[header_row - 1], 1):
        if cell.kind in {"f", "e"} or not isinstance(cell.value, str) or not cell.value.strip():
            raise _problem(name, header_row, column, "表头必须是非空文本，不接受公式、错误值或数字。")
        headers.append(cell.value)
    return SheetSelection(snapshot.sha256, sheet, header_row, tuple(headers))


MONEY_FIELDS = frozenset({"unit_price", "discount_amount", "payment_amount", "refund_amount"})
INTEGER_FIELDS = frozenset({"quantity", "refund_quantity"})


def _convert(cell: CellValue, field: str) -> str:
    value = cell.value
    if cell.kind in {"f", "e"}:
        raise ValueError("不接受公式或 Excel 错误值，不能使用缓存结果。")
    if value is None:
        return ""  # 交给原校验器报告必需值为空，绝不补零。
    if isinstance(value, bool):
        raise ValueError("布尔值不能当作数量、金额或订单文本。")
    if isinstance(value, str):
        return value  # 不去空白、不翻译状态、不改变文本金额精度。
    if isinstance(value, datetime) and field == "order_time":
        if cell.epoch_1900 and cell.number_text is not None and Decimal(60) <= Decimal(cell.number_text) < Decimal(61):
            raise ValueError("Excel 1900 日期系统的虚构 2 月 29 日无法可靠转换。")
        if value.tzinfo is not None or value.microsecond or is_datetime(cell.number_format) != "datetime":
            raise ValueError("日期必须具有完整日期和时间、无小数秒；不会补造时间或猜测时区。")
        return f"{value.year:04d}-{value.month:02d}-{value.day:02d} {value.hour:02d}:{value.minute:02d}:{value.second:02d}"
    if type(value) in (int, float) and field in MONEY_FIELDS | INTEGER_FIELDS:
        number = Decimal(cell.number_text if cell.number_text is not None else str(value))
        digits = number.as_tuple()
        if not number.is_finite() or len(digits.digits) + max(digits.exponent, 0) > 15:
            raise ValueError("Excel 数值非有限或超出 15 位可靠数字范围；应核实后使用精确文本。")
        if field in INTEGER_FIELDS:
            if number != number.to_integral_value():
                raise ValueError("数量必须为整数，不进行四舍五入。")
            return str(int(number))
        # 使用十进制文本补足小数位，不做 float 累加，也不受 Decimal 默认精度舍入。
        text = format(number, "f")
        whole, _, fraction = text.partition(".")
        if len(fraction.rstrip("0")) > 2:
            raise ValueError("金额精度超过两位小数，拒绝四舍五入。")
        return whole + "." + fraction[:2].ljust(2, "0")
    raise ValueError("数据类型无法可靠转换；编号须为文本，日期须为完整时间或标准文本。")


def to_raw_rows(selection: SheetSelection, plan: MappingPlan) -> list[RawCsvRow]:
    """确认是正常交互的前置条件，不是防范恶意 Python 调用者的安全边界。"""
    sheet = selection.sheet
    if not plan.confirmed or plan.context != selection.context or plan.headers != selection.headers:
        raise _problem(sheet.name, selection.header_row, None, "映射未确认或不属于当前快照、工作表和表头。")
    # 防止调用方在预览后意外改坏映射；仍检查完整性和金额含义确认。
    try:
        resolve_mapping(selection.headers, dict(enumerate(plan.assignments, 1)),
                        confirmed_meanings=plan.confirmed_meanings, context=selection.context)
        if len(plan.assignments) != len(selection.headers):
            raise MappingError("映射列数与表头不一致。")
    except MappingError as error:
        raise _problem(sheet.name, selection.header_row, None, str(error)) from error
    rows = []
    for row_number, cells in enumerate(sheet.rows[selection.header_row:], selection.header_row + 1):
        # 即使是明确忽略的额外列，也不能隐藏公式或错误单元格。
        for column, cell in enumerate(cells, 1):
            if cell.kind in {"f", "e"}:
                raise _problem(sheet.name, row_number, column, "存在公式或 Excel 错误值，拒绝使用缓存或忽略。")
        if all(cell.value is None for cell in cells):
            raise _problem(sheet.name, row_number, 1, "空白行不能静默跳过。")
        values = {}
        for column, field in enumerate(plan.assignments, 1):
            if field is None:
                continue
            try:
                values[field] = _convert(cells[column - 1], field)
            except ValueError as error:
                raise _problem(sheet.name, row_number, column, f"[{field}] {error}") from error
        rows.append(RawCsvRow(row_number, row_number, values))
    return rows


def validate_selection(selection: SheetSelection, plan: MappingPlan) -> list[OrderLine]:
    """只在整份数据通过既有 validate_rows 后返回；不计算指标或生成报表。"""
    rows = to_raw_rows(selection, plan)
    try:
        return validate_rows(rows)
    except CsvValidationError as error:
        columns = {field: column for column, field in enumerate(plan.assignments, 1) if field is not None}
        messages = [str(_problem(selection.sheet.name, issue.line_number, columns.get(issue.field),
                                 f"[{issue.field}] {issue.reason}")) for issue in error.issues]
        raise ExcelImportError("\n".join(messages)) from error

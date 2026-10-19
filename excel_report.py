"""Windows Excel 报表：复用完整校验和指标，先核验内容，再无覆盖发布。"""

import io
import os
import tempfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from metrics import MetricsReport, calculate_metrics, format_money
from order_validator import BEIJING_TIME


SHEET_NAMES = ("运营总览", "商品销量排行", "数据校验摘要")
MONEY_FORMAT = '#,##0.00'
COUNT_FORMAT = '#,##0'
DEFAULT_OUTPUT_DIR = Path(__file__).parent / "output"


class ExcelReportError(Exception):
    """导出未完成；不能据此认定已经生成正式报表。"""


def _text(value: str) -> str:
    # Excel 的文本长度有限；openpyxl 默认会截断，必须先拒绝，避免静默丢字。
    if len(value.encode("utf-16-le")) // 2 > 32767:
        raise ExcelReportError("文本超过 Excel 单元格长度限制，拒绝截断导出。")
    return value


def _count(value: int) -> int:
    if len(str(abs(value))) > 15:
        raise ExcelReportError("件数或订单数超过 Excel 的安全数值精度，拒绝导出。")
    return value


def _money(cents) -> Decimal | str:
    if cents is None:
        return format_money(None)
    # 只调用已有展示舍入，不重新计算客单价或任何业务指标。
    value = Decimal(format_money(cents))
    if len(value.as_tuple().digits) > 15:
        raise ExcelReportError("金额（含两位小数）超过 Excel 的安全数值精度，拒绝导出。")
    return value


def _append(sheet, values):
    values = tuple(_text(value) if isinstance(value, str) else value for value in values)
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            # '='、错误码、URL 等一律保存为普通文本；不生成公式或自动链接。
            cell.data_type = "s"
            cell.hyperlink = None


def _style(sheet, widths):
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="244A64")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 42
        for cell in row:
            cell.font = Font(name="微软雅黑", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True,
                                       horizontal="right" if cell.data_type == "n" else "left")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="EFF5F8")
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def _build_workbook(report: MetricsReport, source_name: str, generated_at: datetime) -> Workbook:
    workbook = Workbook()
    overview = workbook.active
    overview.title = SHEET_NAMES[0]
    _append(overview, ("指标", "数值", "单位", "统计口径"))
    rows = (
        ("商品明细数", _count(report.line_count), "条", "一行是一条商品明细，不等于一笔订单"),
        ("唯一订单数", _count(report.total_order_count), "笔", "包含未支付和取消订单；按订单编号去重"),
        ("支付金额", _money(report.payment_cents), "元", "包含后来全额、部分退款订单的历史支付"),
        ("退款金额", _money(report.refund_cents), "元", "已完成的退款金额"),
        ("退款后净销售额", _money(report.net_sales_cents), "元", "支付金额 − 退款金额；不等于利润"),
        ("支付订单数", _count(report.paid_order_count), "笔", "包含后来全额、部分退款的唯一已支付订单"),
        ("净有效订单数", _count(report.net_order_count), "笔", "已支付订单中，整单净金额大于零"),
        ("支付客单价", _money(report.paid_average_cents), "元/笔", "支付金额 ÷ 支付订单数；仅展示时四舍五入"),
        ("净客单价", _money(report.net_average_cents), "元/笔", "净销售额 ÷ 净有效订单数；仅展示时四舍五入"),
    )
    for values in rows:
        _append(overview, values)
        overview.cell(overview.max_row, 2).number_format = MONEY_FORMAT if values[2].startswith("元") else COUNT_FORMAT
    _style(overview, {"A": 24, "B": 23, "C": 12, "D": 64})

    ranking = workbook.create_sheet(SHEET_NAMES[1])
    _append(ranking, ("排名", "商品编号", "商品名称", "商品类别", "净成交件数"))
    for rank, product in enumerate(report.products, 1):
        _append(ranking, (_count(rank), _text(product.product_id), _text(product.product_name),
                          _text(product.category), _count(product.net_quantity)))
        ranking.cell(ranking.max_row, 1).number_format = COUNT_FORMAT
        ranking.cell(ranking.max_row, 5).number_format = COUNT_FORMAT
    ranking.auto_filter.ref = f"A1:E{ranking.max_row}"
    _style(ranking, {"A": 10, "B": 18, "C": 34, "D": 20, "E": 18})

    summary = workbook.create_sheet(SHEET_NAMES[2])
    _append(summary, ("项目", "内容"))
    source = report.source
    is_excel = source is not None and source.file_format == "XLSX"
    for values in (
        ("输入 Excel 文件名" if is_excel else "输入 CSV 文件名", _text(source.file_name if source else source_name)),
        ("校验结果", "整份数据通过既有严格校验；未丢弃异常行或自动补值"),
        ("商品明细数", _count(report.line_count)),
        ("唯一订单数", _count(report.total_order_count)),
        ("金额口径", "人民币元；历史支付包含全额退款订单；净销售额 = 支付 − 已完成退款"),
        ("订单与客单价口径", "支付与净额口径分别使用对应订单数；零分母显示无法计算"),
        ("商品排行口径", "购买件数 − 明确退款件数；只退差价不扣件数；并列按商品编号字符串升序"),
        ("数据时间边界", "订单状态快照；缺少支付和退款发生时间，不是每日资金流水"),
        ("报表生成时间（北京时间 UTC+08:00）", generated_at),
        ("使用提示", "请使用虚构模拟订单练习。修改 Excel 后需重新读取并确认映射，本报表不会自动更新" if is_excel else
         "本项目附带 CSV 为虚构模拟数据。修改 CSV 后需重新运行程序，本报表不会自动更新"),
    ):
        _append(summary, values)
    # 追加来源行，不移动原有指标、生成时间及其格式位置，保留 CSV 报表兼容性。
    _append(summary, ("输入文件格式", source.file_format if source else "CSV"))
    if is_excel:
        for values in (("输入工作表", source.sheet_name), ("表头原始行号", source.header_row),
                       ("输入快照 SHA256", source.sha256),
                       ("快照说明", "结果对应映射确认时的内存快照；源文件后来变化需重新运行并确认")):
            _append(summary, values)
            if values[0] == "表头原始行号":
                summary.cell(summary.max_row, 2).number_format = COUNT_FORMAT
    summary["B4"].number_format = COUNT_FORMAT
    summary["B5"].number_format = COUNT_FORMAT
    summary["B10"].number_format = "yyyy-mm-dd hh:mm:ss"
    _style(summary, {"A": 44, "B": 100})
    # 文字较长的两列保持足够高度；不靠缩小字体挤入单元格。
    for index in range(6, 10):
        summary.row_dimensions[index].height = 54
    return workbook


def _serialize_and_verify(workbook: Workbook) -> bytes:
    """先在内存完成 XLSX，并回读检查每个值和关键格式，不产生正式文件。"""
    with io.BytesIO() as buffer:
        workbook.save(buffer)
        buffer.seek(0)
        restored = load_workbook(buffer, data_only=False)
        try:
            if restored.sheetnames != list(SHEET_NAMES):
                raise ExcelReportError("生成的工作表清单异常。")
            for original, saved in zip(workbook.worksheets, restored.worksheets):
                if (original.max_row, original.max_column, original.freeze_panes, original.auto_filter.ref) != (
                        saved.max_row, saved.max_column, saved.freeze_panes, saved.auto_filter.ref):
                    raise ExcelReportError("工作表结构或表头设置回读不一致。")
                for row in original:
                    for expected in row:
                        actual = saved[expected.coordinate]
                        if actual.data_type == "f" or actual.hyperlink is not None:
                            raise ExcelReportError("发现非预期公式或链接，拒绝导出。")
                        if isinstance(expected.value, Decimal):
                            equal = actual.data_type == "n" and Decimal(str(actual.value)) == expected.value
                        else:
                            equal = actual.value == expected.value and actual.data_type == expected.data_type
                        if not equal or actual.number_format != expected.number_format:
                            raise ExcelReportError(f"{saved.title}!{actual.coordinate} 的内容或数值精度回读不一致。")
        finally:
            restored.close()
        return buffer.getvalue()


def _write_payload(stream, payload: bytes):
    if stream.write(payload) != len(payload):
        raise OSError("报表临时文件未完整写入。")
    stream.flush()
    os.fsync(stream.fileno())


def _publish_no_replace(temporary: Path, destination: Path):
    # Windows 的 rename 在目标存在时拒绝执行；不能用会覆盖目标的 replace。
    if os.name != "nt":
        raise ExcelReportError("当前无覆盖发布流程仅支持 Windows；其他系统尚未实现。")
    os.rename(temporary, destination)


def export_excel(csv_path: str | Path, output_path: str | Path | None = None) -> Path:
    """完整校验 → 复用指标 → 内存回读 → 临时文件 → Windows 无覆盖发布。"""
    report = calculate_metrics(csv_path)  # 校验异常直接向上传递；此前不创建输出目录。
    generated_at = datetime.now(BEIJING_TIME).replace(tzinfo=None, microsecond=0)
    destination = Path(output_path) if output_path is not None else (
        DEFAULT_OUTPUT_DIR / f"运营报表_{generated_at:%Y%m%d_%H%M%S}_{uuid4().hex[:8]}.xlsx"
    )
    temporary = None
    workbook = None
    try:
        # 在创建文件前确定路径，正式发布后不再做可能失败的路径解析。
        destination = destination.absolute()
        if destination.suffix.lower() != ".xlsx":
            raise ExcelReportError("输出文件必须使用 .xlsx 扩展名。")
        if os.path.lexists(destination):
            raise ExcelReportError(f"目标已经存在，拒绝覆盖：{destination}")
        workbook = _build_workbook(report, Path(csv_path).name, generated_at)
        payload = _serialize_and_verify(workbook)
        workbook.close()
        workbook = None
        destination.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(mode="wb", prefix=".excel_report_", suffix=".tmp",
                                         dir=destination.parent, delete=False) as stream:
            temporary = Path(stream.name)
            _write_payload(stream, payload)
        _publish_no_replace(temporary, destination)
        temporary = None
        return destination
    except ExcelReportError:
        raise
    except Exception as error:
        # 只在导出边界包装异常，不返回成功；原始原因保留在异常链中。
        raise ExcelReportError(f"导出失败，未发布本次正式报表：{error}") from error
    finally:
        if workbook is not None:
            workbook.close()
        if temporary is not None:
            try:
                temporary.unlink(missing_ok=True)
            except OSError as error:
                raise ExcelReportError(f"导出未完成，临时文件无法清理，请检查：{temporary}") from error

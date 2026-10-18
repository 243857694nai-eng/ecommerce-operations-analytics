"""Excel 5A 隔离测试：只在系统临时目录创建模拟工作簿，结束后清理。"""

import contextlib
import hashlib
import io
import subprocess
import sys
import tempfile
import unittest
import zipfile
from dataclasses import replace
from datetime import date, datetime, time, timedelta
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from csv_reader import REQUIRED_FIELDS, read_csv
from excel_reader import (ExcelImportError, header_candidates, read_snapshot, select_sheet,
                          to_raw_rows, validate_selection)
from field_mapping import ALIASES, confirm_mapping, resolve_mapping
from inspect_excel import main
from order_validator import load_and_validate
from test_csv_validation import sample_row, second_row


ROOT = Path(__file__).parent


class ExcelReaderTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="excel_import_test_")
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)
        self.path = self.folder / "模拟订单.xlsx"

    def write(self, rows=None, headers=REQUIRED_FIELDS, header_row=1, extra_sheet=False, edit=None, epoch=None,
              iso_dates=False):
        workbook = Workbook()
        workbook.iso_dates = iso_dates
        sheet = workbook.active
        sheet.title = "模拟订单"
        if epoch:
            workbook.epoch = epoch
        for _ in range(header_row - 1):
            sheet.append(["模拟说明，不作为交易数据"])
        sheet.append(headers)
        for row in ([sample_row(), second_row()] if rows is None else rows):
            sheet.append([row.get(field, "") for field in headers])
        if extra_sheet:
            workbook.create_sheet("其他页").append(["不是订单数据"])
            workbook.active = 1
        if edit:
            edit(sheet)
        workbook.save(self.path)
        workbook.close()
        return self.path

    def selected(self, header_row=1, choices=None, meanings=()):
        snapshot = read_snapshot(self.path)
        selection = select_sheet(snapshot, "模拟订单", header_row)
        plan = resolve_mapping(selection.headers, choices, confirmed_meanings=meanings, context=selection.context)
        return selection, confirm_mapping(plan, True)

    def rejected(self, contains=None, **kwargs):
        before = hashlib.sha256(self.path.read_bytes()).digest()
        with self.assertRaises(ExcelImportError) as caught:
            selection, plan = self.selected(**kwargs)
            validate_selection(selection, plan)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).digest(), before)
        self.assertFalse(hasattr(caught.exception, "records"))
        if contains:
            self.assertIn(contains, str(caught.exception))
        return str(caught.exception)

    def run_cli(self, answers):
        output = io.StringIO()
        with contextlib.redirect_stdout(output), patch("builtins.input", side_effect=answers):
            code = main([str(self.path)])
        return code, output.getvalue()

    def change_xml(self, part, old, new):
        with zipfile.ZipFile(io.BytesIO(self.path.read_bytes())) as archive:
            contents = {name: archive.read(name) for name in archive.namelist()}
        self.assertIn(old, contents[part])
        contents[part] = contents[part].replace(old, new)
        with zipfile.ZipFile(self.path, "w") as archive:
            for name, payload in contents.items():
                archive.writestr(name, payload)

    def write_raw_iso_date(self, text, *, header_row=1, last_row=False):
        """直接修改临时 XLSX 的 t=d 日期原文，不能只模拟已解析的 datetime。"""
        moment = datetime(2026, 10, 1, 10, 7, 11)
        self.write([sample_row(order_time=moment), second_row(order_time=moment)],
                   header_row=header_row, iso_dates=True)
        with zipfile.ZipFile(io.BytesIO(self.path.read_bytes())) as archive:
            contents = {name: archive.read(name) for name in archive.namelist()}
        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        root = ElementTree.fromstring(contents["xl/worksheets/sheet1.xml"])
        coordinate = f"C{header_row + (2 if last_row else 1)}"
        cell = root.find(f".//{namespace}c[@r='{coordinate}']")
        self.assertEqual(cell.attrib["t"], "d")
        cell.find(namespace + "v").text = text
        contents["xl/worksheets/sheet1.xml"] = ElementTree.tostring(root, encoding="utf-8")
        with zipfile.ZipFile(self.path, "w") as archive:
            for name, payload in contents.items():
                archive.writestr(name, payload)
        # 重新打开原始 XML，证明测试包含的确实是目标原文。
        with zipfile.ZipFile(self.path) as archive:
            saved = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        self.assertEqual(saved.find(f".//{namespace}c[@r='{coordinate}']/{namespace}v").text, text)

    def test_native_iso_utc_suffix_is_rejected_from_original_xml(self):
        self.write_raw_iso_date("2026-10-01T10:07:11Z")
        self.rejected("第 2 行 C 列")

    def test_native_iso_timezone_offset_is_rejected_from_original_xml(self):
        self.write_raw_iso_date("2026-10-01T10:07:11+09:00")
        self.rejected("第 2 行 C 列")

    def test_native_iso_invalid_suffix_is_rejected_from_original_xml(self):
        self.write_raw_iso_date("2026-10-01T10:07:11INVALID")
        self.rejected("第 2 行 C 列")

    def test_valid_native_iso_datetime_and_zero_fraction_remain_supported(self):
        for text in ("2026-10-01T10:07:11", "2026-10-01T10:07:11.000", "2026-10-01T10:07:11.0000000"):
            with self.subTest(text=text):
                self.write_raw_iso_date(text)
                before = hashlib.sha256(self.path.read_bytes()).digest()
                selection, plan = self.selected()
                self.assertEqual(to_raw_rows(selection, plan)[0].values["order_time"], "2026-10-01 10:07:11")
                self.assertEqual(len(validate_selection(selection, plan)), 2)
                self.assertEqual(hashlib.sha256(self.path.read_bytes()).digest(), before)

    def test_raw_iso_check_happens_before_openpyxl_can_discard_information(self):
        self.write_raw_iso_date("2026-10-01T10:07:11Z")
        with patch("excel_reader.load_workbook") as parser:
            self.rejected("原生 ISO 日期")
        parser.assert_not_called()

    def test_native_iso_impossible_calendar_date_reports_original_cell(self):
        for text in ("2026-02-30T10:07:11", "2026-13-01T10:07:11", "2026-10-01T24:00:00"):
            with self.subTest(text=text):
                self.write_raw_iso_date(text)
                self.rejected("工作表 '模拟订单' 第 2 行 C 列")

    def test_native_iso_nonzero_fraction_is_not_truncated_or_rounded(self):
        for suffix in (".123", ".0001", ".0000001", ".1234567"):
            with self.subTest(suffix=suffix):
                self.write_raw_iso_date("2026-10-01T10:07:11" + suffix)
                self.rejected("第 2 行 C 列")

    def test_native_iso_other_offsets_spaces_and_partial_timestamps_are_rejected(self):
        for text in ("2026-10-01T10:07:11+00:00", "2026-10-01T10:07:11-05:00",
                     "2026-10-01T10:07:11z", "2026-10-01T10:07:11 ",
                     "2026-10-01T10:07", "2026-10-01T10:07:11\n"):
            with self.subTest(text=text):
                self.write_raw_iso_date(text)
                self.rejected("第 2 行 C 列")

    def test_invalid_iso_on_last_row_refuses_whole_file_without_calling_validator(self):
        self.write_raw_iso_date("2026-10-01T10:07:11Z", header_row=4, last_row=True)
        before = hashlib.sha256(self.path.read_bytes()).digest()
        with patch("excel_reader.validate_rows") as validator:
            code, output = self.run_cli([])
        self.assertEqual(code, 1)
        self.assertIn("工作表 '模拟订单' 第 6 行 C 列", output)
        self.assertNotIn("整份校验通过", output)
        self.assertNotIn("2026-10-01T10:07:11Z", output)
        validator.assert_not_called()
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).digest(), before)

    def test_raw_iso_date_only_and_time_only_do_not_become_complete_order_times(self):
        for text in ("2026-10-01", "10:07:11"):
            with self.subTest(text=text):
                self.write_raw_iso_date(text)
                self.rejected("第 2 行 C 列")

    def test_native_iso_regression_runs_actual_cli_without_success_or_partial_output(self):
        for text in ("2026-10-01T10:07:11Z", "2026-10-01T10:07:11+09:00",
                     "2026-10-01T10:07:11INVALID"):
            with self.subTest(text=text):
                self.write_raw_iso_date(text)
                before = hashlib.sha256(self.path.read_bytes()).digest()
                result = subprocess.run([sys.executable, "-B", "-X", "utf8", str(ROOT / "inspect_excel.py"), str(self.path)],
                                        input="", cwd=self.folder, capture_output=True, text=True, encoding="utf-8",
                                        check=False, timeout=30)
                self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
                self.assertIn("工作表 '模拟订单' 第 2 行 C 列", result.stdout)
                self.assertNotIn("整份校验通过", result.stdout)
                self.assertEqual(hashlib.sha256(self.path.read_bytes()).digest(), before)

    def test_normal_file_reuses_existing_validator_and_is_readonly(self):
        self.write()
        before = self.path.read_bytes()
        selection, plan = self.selected()
        with patch("excel_reader.validate_rows", wraps=__import__("order_validator").validate_rows) as validator:
            records = validate_selection(selection, plan)
        validator.assert_called_once()
        self.assertEqual(len(records), 2)
        self.assertEqual(records[0].payment_amount_cents, 3582)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(list(self.folder.iterdir()), [self.path])

    def test_sample_excel_and_csv_produce_identical_raw_rows_and_records(self):
        csv_path = ROOT / "data" / "sample_orders.csv"
        before = hashlib.sha256(csv_path.read_bytes()).digest()
        csv_rows = read_csv(csv_path)
        self.write([row.values for row in csv_rows])
        selection, plan = self.selected()
        self.assertEqual(to_raw_rows(selection, plan), csv_rows)
        self.assertEqual(validate_selection(selection, plan), load_and_validate(csv_path))
        self.assertEqual(hashlib.sha256(csv_path.read_bytes()).digest(), before)

    def test_chinese_headers_and_reordered_columns_are_supported(self):
        def edit(sheet):
            for column, field in enumerate(reversed(REQUIRED_FIELDS), 1):
                sheet.cell(1, column).value = ALIASES[field][0]
        self.write(headers=tuple(reversed(REQUIRED_FIELDS)), edit=edit)
        selection, plan = self.selected()
        self.assertEqual(validate_selection(selection, plan)[0].order_id, "A001")

    def test_explicit_sheet_selection_does_not_use_active_sheet(self):
        self.write(extra_sheet=True)
        snapshot = read_snapshot(self.path)
        self.assertEqual([sheet.name for sheet in snapshot.sheets], ["模拟订单", "其他页"])
        selection, plan = self.selected()
        self.assertEqual(len(validate_selection(selection, plan)), 2)

    def test_header_candidates_are_advisory_and_original_row_numbers_survive(self):
        self.write(header_row=4)
        snapshot = read_snapshot(self.path)
        self.assertIn((4, 14), header_candidates(snapshot.sheets[0]))
        selection, plan = self.selected(header_row=4)
        self.assertEqual([row.source_line for row in validate_selection(selection, plan)], [5, 6])

    def test_wrong_sheet_and_header_positions_have_no_fallback(self):
        self.write()
        snapshot = read_snapshot(self.path)
        for name, row in (("不存在", 1), ("模拟订单", 0), ("模拟订单", 100), ("模拟订单", True)):
            with self.subTest(name=name, row=row), self.assertRaises(ExcelImportError):
                select_sheet(snapshot, name, row)

    def test_duplicate_headers_keep_both_columns_and_require_explicit_choice(self):
        self.write(headers=REQUIRED_FIELDS + ("order_id",))
        snapshot = read_snapshot(self.path)
        selection = select_sheet(snapshot, "模拟订单", 1)
        self.assertEqual(len(selection.headers), 15)
        plan = confirm_mapping(resolve_mapping(selection.headers, {1: "order_id", 15: None},
                                               context=selection.context), True)
        self.assertEqual(len(validate_selection(selection, plan)), 2)

    def test_unconfirmed_plan_never_calls_validator(self):
        self.write()
        selection, plan = self.selected()
        with patch("excel_reader.validate_rows") as validator, self.assertRaisesRegex(ExcelImportError, "未确认"):
            validate_selection(selection, replace(plan, confirmed=False))
        validator.assert_not_called()

    def test_plan_bound_to_snapshot_sheet_and_header(self):
        self.write(extra_sheet=True)
        selection, plan = self.selected()
        for changed in (replace(plan, context="another snapshot"), replace(plan, headers=tuple(reversed(plan.headers)))):
            with self.subTest(plan=changed), self.assertRaisesRegex(ExcelImportError, "不属于"):
                validate_selection(selection, changed)

    def test_source_replacement_after_preview_does_not_change_memory_snapshot(self):
        self.write()
        selection, plan = self.selected()
        # 模拟用户用另一版本覆盖自己的临时文件；读取器不重新打开该路径。
        self.write([sample_row(order_status="未知状态")])
        replacement = self.path.read_bytes()
        self.assertEqual(len(validate_selection(selection, plan)), 2)
        self.assertEqual(self.path.read_bytes(), replacement)
        self.rejected("order_status")

    def test_modified_mapping_cannot_omit_or_duplicate_required_fields(self):
        self.write()
        selection, plan = self.selected()
        for assignments in (plan.assignments[:-1], ("order_id",) * 14, plan.assignments + (None,)):
            with self.subTest(assignments=assignments), self.assertRaises(ExcelImportError):
                validate_selection(selection, replace(plan, assignments=assignments))

    def test_text_identifiers_preserve_leading_zeros_and_long_digits(self):
        long_id = "0" + "1234567890" * 4
        self.write([sample_row(order_id="00001", order_line_id=long_id, product_id="00009")])
        selection, plan = self.selected()
        record = validate_selection(selection, plan)[0]
        self.assertEqual((record.order_id, record.order_line_id, record.product_id), ("00001", long_id, "00009"))

    def test_numeric_identifier_is_rejected_even_with_zero_padded_display(self):
        def edit(sheet):
            sheet["A2"] = 123
            sheet["A2"].number_format = "000000"
        self.write(edit=edit)
        self.rejected("第 2 行 A 列")

    def test_numeric_integer_quantities_and_cents_convert_exactly(self):
        self.write([sample_row(quantity=2, refund_quantity=0, unit_price=19.9, discount_amount=3.98, payment_amount=35.82,
                               refund_amount=0)])
        selection, plan = self.selected()
        raw = to_raw_rows(selection, plan)[0]
        self.assertEqual(raw.values["unit_price"], "19.90")
        self.assertEqual(raw.values["refund_amount"], "0.00")
        self.assertEqual(validate_selection(selection, plan)[0].payment_amount_cents, 3582)

    def test_text_money_keeps_existing_strict_format(self):
        self.write([sample_row(unit_price="19.9")])
        self.rejected("恰有两位小数")

    def test_subcent_numeric_money_is_rejected_instead_of_rounded(self):
        for amount in (1.005, 0.001, 19.9001):
            with self.subTest(amount=amount):
                self.write([sample_row(unit_price=amount)])
                self.rejected("金额精度超过")

    def test_numeric_amount_over_reliable_precision_is_rejected(self):
        for amount in (1234567890123456, 1e20):
            with self.subTest(amount=amount):
                self.write([sample_row(unit_price=amount)])
                self.rejected("15 位")

    def test_numeric_xml_precision_is_checked_before_float_can_erase_it(self):
        self.write([sample_row(unit_price=19.9)])
        self.change_xml("xl/worksheets/sheet1.xml", b"<v>19.9</v>", b"<v>19.90000000000000001</v>")
        self.rejected("第 2 行 H 列")

    def test_long_text_amount_remains_exact_and_is_not_float_limited(self):
        huge = "123456789012345678901234567890.12"
        self.write([sample_row(quantity="1", unit_price=huge, discount_amount="0.00", payment_amount=huge)])
        selection, plan = self.selected()
        self.assertEqual(validate_selection(selection, plan)[0].unit_price_cents, int(huge.replace(".", "")))

    def test_fractional_and_boolean_quantities_are_rejected(self):
        for quantity in (1.5, True, False):
            with self.subTest(quantity=quantity):
                self.write([sample_row(quantity=quantity)])
                self.rejected("第 2 行 G 列")

    def test_illegal_integer_quantities_use_existing_rules(self):
        for quantity in (0, -1, "01", "1.0"):
            with self.subTest(quantity=quantity):
                self.write([sample_row(quantity=quantity)])
                self.rejected("正整数")

    def test_native_datetime_and_both_excel_date_epochs(self):
        for epoch in (None, CALENDAR_MAC_1904):
            with self.subTest(epoch=epoch):
                self.write([sample_row(order_time=datetime(2026, 10, 1, 10, 7, 11))], epoch=epoch)
                selection, plan = self.selected()
                self.assertEqual(to_raw_rows(selection, plan)[0].values["order_time"], "2026-10-01 10:07:11")
                self.assertEqual(validate_selection(selection, plan)[0].order_time.utcoffset(), timedelta(hours=8))

    def test_date_only_time_only_fractional_seconds_and_serial_without_style_rejected(self):
        for moment in (date(2026, 10, 1), time(10, 7), datetime(2026, 10, 1, 10, 7, 11, 123000), 46300):
            with self.subTest(moment=moment):
                self.write([sample_row(order_time=moment)])
                self.rejected("第 2 行 C 列")

    def test_date_text_is_not_stripped_or_guessed(self):
        self.write([sample_row(order_time="2026/10/01")])
        self.rejected("日期格式")

    def test_excel_fictitious_1900_leap_day_is_not_silently_mapped_to_february_28(self):
        self.write([sample_row(order_time=datetime(1900, 2, 28, 0, 0, 0))])
        self.change_xml("xl/worksheets/sheet1.xml", b"<v>59</v>", b"<v>60</v>")
        self.rejected("虚构 2 月 29 日")

    def test_unknown_status_is_not_translated(self):
        self.write([sample_row(order_status="交易成功")])
        self.rejected("第 2 行 K 列")

    def test_missing_refund_quantity_is_not_filled_with_zero(self):
        self.write([sample_row(refund_quantity=None)])
        self.rejected("第 2 行 M 列")

    def test_refund_limits_and_whole_order_consistency_are_reused(self):
        cases = ([sample_row(refund_amount="99.00")], [sample_row(refund_quantity="3")],
                 [sample_row(), second_row(channel="直播间")], [sample_row(), second_row(order_line_id="A001-01")])
        for rows in cases:
            with self.subTest(rows=rows):
                self.write(rows)
                self.rejected()

    def test_late_invalid_row_returns_no_partial_records(self):
        self.write([sample_row(), second_row(order_status="未知")])
        self.assertIn("第 3 行 K 列", self.rejected())

    def test_formula_cell_is_rejected_with_original_cell_address(self):
        self.write([sample_row(payment_amount="=2*19.90-3.98")])
        self.rejected("第 2 行 J 列")

    def test_formula_with_numeric_cached_result_is_not_treated_as_payment(self):
        self.write([sample_row(payment_amount="=2*19.90-3.98")])
        # 在隔离测试副本中模拟办公软件保存的公式缓存，仍须识别为公式。
        with zipfile.ZipFile(io.BytesIO(self.path.read_bytes())) as archive:
            contents = {name: archive.read(name) for name in archive.namelist()}
        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        root = ElementTree.fromstring(contents["xl/worksheets/sheet1.xml"])
        cell = root.find(f".//{namespace}c[@r='J2']")
        self.assertIsNotNone(cell.find(namespace + "f"))
        cell.find(namespace + "v").text = "35.82"
        contents["xl/worksheets/sheet1.xml"] = ElementTree.tostring(root, encoding="utf-8")
        with zipfile.ZipFile(self.path, "w") as archive:
            for name, payload in contents.items():
                archive.writestr(name, payload)
        self.rejected("第 2 行 J 列")

    def test_formula_in_explicitly_ignored_column_is_still_rejected(self):
        self.write(headers=REQUIRED_FIELDS + ("备注",), edit=lambda sheet: setattr(sheet["O2"], "value", "=1+1"))
        self.rejected("第 2 行 O 列", choices={15: None})

    def test_error_cell_is_rejected_without_printing_original_value(self):
        self.write([sample_row(product_name="#N/A")])
        error = self.rejected("第 2 行 E 列")
        self.assertNotIn("#N/A", error)

    def test_literal_text_starting_equals_is_not_a_formula(self):
        def edit(sheet):
            sheet["E2"] = "=模拟文本"
            sheet["E2"].data_type = "s"
        self.write([sample_row()], edit=edit)
        selection, plan = self.selected()
        self.assertEqual(validate_selection(selection, plan)[0].product_name, "=模拟文本")

    def test_merged_cells_are_rejected(self):
        self.write(edit=lambda sheet: sheet.merge_cells("A2:B2"))
        self.rejected("A2:B2")

    def test_formula_numeric_and_blank_headers_report_column(self):
        for value in ("=1+1", 123, None):
            with self.subTest(value=value):
                self.write(edit=lambda sheet: setattr(sheet["B1"], "value", value))
                self.rejected("第 1 行 B 列")

    def test_blank_middle_row_is_not_skipped(self):
        self.write([sample_row(), {}, second_row()])
        self.rejected("第 3 行 A 列")

    def test_totals_row_is_not_skipped(self):
        self.write([sample_row(), {"order_id": "合计", "quantity": 2}])
        self.rejected("第 3 行")

    def test_hidden_rows_and_columns_are_not_silently_discarded(self):
        for edit, location in ((lambda sheet: setattr(sheet.row_dimensions[2], "hidden", True), "第 2 行"),
                               (lambda sheet: setattr(sheet.column_dimensions["G"], "hidden", True), "G 列")):
            with self.subTest(location=location):
                self.write(edit=edit)
                self.rejected(location)

    def test_header_only_file_is_not_a_successful_empty_dataset(self):
        self.write([])
        self.rejected("没有可校验")

    def test_hidden_sheet_is_rejected_but_other_visible_sheet_is_not_implicitly_selected(self):
        self.write(extra_sheet=True, edit=lambda sheet: setattr(sheet, "sheet_state", "hidden"))
        self.rejected("工作表被隐藏")

    def test_missing_required_header_reports_sheet_and_header_in_cli(self):
        self.write(headers=tuple(field for field in REQUIRED_FIELDS if field != "refund_quantity"))
        code, output = self.run_cli(["1", "1"] + [""] * 13)
        self.assertEqual(code, 1)
        self.assertIn("工作表 '模拟订单' 第 1 行表头", output)
        self.assertIn("缺少必需字段：refund_quantity", output)

    def test_extra_cell_beyond_header_is_not_silently_dropped(self):
        self.write(edit=lambda sheet: setattr(sheet["O3"], "value", "未说明的额外数据"))
        self.rejected("第 1 行 O 列")

    def test_formatted_blank_trailing_row_is_not_silently_dropped(self):
        self.write(edit=lambda sheet: setattr(sheet["A4"], "number_format", "@"))
        self.rejected("第 4 行 A 列")

    def test_damaged_and_encrypted_container_files_are_rejected_readonly(self):
        for payload, message in ((b"not xlsx", "无法读取"), (bytes.fromhex("D0CF11E0A1B11AE1") + b"mock", "加密")):
            with self.subTest(message=message):
                self.path.write_bytes(payload)
                self.rejected(message)

    def test_unsupported_extension_and_missing_file_are_errors(self):
        for path in (self.folder / "模拟.xls", self.folder / "不存在.xlsx"):
            with self.subTest(path=path), self.assertRaises(ExcelImportError):
                read_snapshot(path)

    def test_parser_exception_preserves_source_hash_and_hides_data(self):
        self.write()
        before = self.path.read_bytes()
        with patch("excel_reader.load_workbook", side_effect=ValueError("模拟客户机密")):
            self.assertNotIn("模拟客户机密", self.rejected("ValueError"))
        self.assertEqual(self.path.read_bytes(), before)

    def test_unexpected_validator_exception_is_not_reported_as_success(self):
        self.write()
        before = self.path.read_bytes()
        selection, plan = self.selected()
        with patch("excel_reader.validate_rows", side_effect=RuntimeError("模拟故障")), self.assertRaises(RuntimeError):
            validate_selection(selection, plan)
        self.assertEqual(self.path.read_bytes(), before)

    def test_cli_unexpected_exception_reports_failure_without_raw_data(self):
        self.write()
        before = self.path.read_bytes()
        with patch("excel_reader.validate_selection", side_effect=RuntimeError("虚构客户敏感原文")):
            code, output = self.run_cli(["1", "1"] + [""] * 14 + ["确认"])
        self.assertEqual(code, 1)
        self.assertIn("程序异常（RuntimeError）", output)
        self.assertNotIn("虚构客户敏感原文", output)
        self.assertNotIn("整份校验通过", output)
        self.assertEqual(self.path.read_bytes(), before)

    def test_cli_requires_final_confirmation_and_shows_mapping_before_validation(self):
        self.write()
        with patch("excel_reader.validate_selection") as validator:
            code, output = self.run_cli(["1", "1"] + [""] * 14 + ["取消"])
        self.assertEqual(code, 2)
        self.assertIn("最终映射预览", output)
        validator.assert_not_called()
        self.assertNotIn("校验通过", output)

    def test_cli_success_only_prints_validation_summary_and_never_exports(self):
        self.write()
        before = self.path.read_bytes()
        code, output = self.run_cli(["1", "1"] + [""] * 14 + ["确认"])
        self.assertEqual(code, 0)
        self.assertIn("整份校验通过：2 条商品明细，1 个唯一订单编号", output)
        self.assertNotIn("学习文具", output)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(list(self.folder.iterdir()), [self.path])

    def test_cli_cancel_sheet_eof_and_interrupt_do_not_validate(self):
        self.write()
        for answers in (["q"], [EOFError()], [KeyboardInterrupt()]):
            with self.subTest(answers=answers), patch("excel_reader.validate_selection") as validator:
                code, _ = self.run_cli(answers)
                self.assertEqual(code, 2)
                validator.assert_not_called()

    def test_cli_invalid_sheet_never_chooses_default(self):
        self.write()
        code, output = self.run_cli(["99"])
        self.assertEqual(code, 1)
        self.assertIn("序号无效", output)

    def test_cli_ambiguous_amount_requires_explicit_meaning_before_final_confirmation(self):
        self.write(edit=lambda sheet: setattr(sheet["J1"], "value", "实付金额"))
        with patch("excel_reader.validate_selection") as validator:
            code, _ = self.run_cli(["1", "1"] + [""] * 10 + ["不确定"])
        self.assertEqual(code, 2)
        validator.assert_not_called()
        code, output = self.run_cli(["1", "1"] + [""] * 10 + ["明细"] + [""] * 4 + ["确认"])
        self.assertEqual(code, 0, output)

    def test_cli_extra_column_must_be_explicitly_ignored_and_not_print_data(self):
        self.write(headers=REQUIRED_FIELDS + ("备注",), edit=lambda sheet: setattr(sheet["O2"], "value", "虚构客户资料"))
        code, output = self.run_cli(["1", "1"] + [""] * 15)
        self.assertEqual(code, 1)
        self.assertNotIn("虚构客户资料", output)
        code, output = self.run_cli(["1", "1"] + [""] * 14 + ["-", "确认"])
        self.assertEqual(code, 0, output)
        self.assertNotIn("虚构客户资料", output)

    def test_cli_invalid_data_does_not_print_success_or_partial_statistics(self):
        self.write([sample_row(order_status="未知状态")])
        code, output = self.run_cli(["1", "1"] + [""] * 14 + ["确认"])
        self.assertEqual(code, 1)
        self.assertIn("第 2 行 K 列", output)
        self.assertNotIn("整份校验通过", output)

    def test_import_modules_have_no_io_or_validation_side_effects(self):
        script = """
from unittest.mock import patch
with patch('pathlib.Path.read_bytes', side_effect=AssertionError('read')), \\
     patch('pathlib.Path.write_bytes', side_effect=AssertionError('write')), \\
     patch('order_validator.validate_rows', side_effect=AssertionError('validate')), \\
     patch('builtins.input', side_effect=AssertionError('prompt')):
    import field_mapping, excel_reader, inspect_excel
print('safe imports')
"""
        result = subprocess.run([sys.executable, "-B", "-X", "utf8", "-c", script], cwd=ROOT,
                                capture_output=True, text=True, encoding="utf-8", check=False)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("safe imports", result.stdout)

    def test_real_cli_process_with_simulated_stdin_from_another_directory(self):
        self.write()
        before = self.path.read_bytes()
        answers = "\n".join(["1", "1"] + [""] * 14 + ["确认"]) + "\n"
        result = subprocess.run([sys.executable, "-B", "-X", "utf8", str(ROOT / "inspect_excel.py"), str(self.path)],
                                input=answers, cwd=self.folder, capture_output=True, text=True, encoding="utf-8",
                                check=False)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        self.assertIn("整份校验通过：2 条商品明细", result.stdout)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(list(self.folder.iterdir()), [self.path])


if __name__ == "__main__":
    unittest.main()

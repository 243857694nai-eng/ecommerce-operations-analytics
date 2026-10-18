"""Excel 导出集成与失败测试：只在新建的隔离临时目录写文件。"""

import contextlib
import csv
import hashlib
import io
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

import excel_report
from csv_reader import CsvValidationError, REQUIRED_FIELDS
from excel_report import ExcelReportError, MONEY_FORMAT, SHEET_NAMES, export_excel
from metrics import calculate_metrics, format_money
from order_validator import BEIJING_TIME
from report_excel import main


ROOT = Path(__file__).parent
SAMPLE = ROOT / "data" / "sample_orders.csv"


def row(order_id="A1", product_id="P1", **changes):
    values = dict(zip(REQUIRED_FIELDS, (
        order_id, order_id + "-01", "2026-10-01 10:00:00", product_id,
        "模拟商品" + product_id, "文具用品", "1", "10.00", "0.00", "10.00",
        "已支付", "0.00", "0", "平台搜索",
    )))
    values.update(changes)
    return values


class ExcelReportTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="excel_report_test_", dir=ROOT)
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)
        self.source = self.folder / "模拟订单.csv"
        self.source.write_bytes(SAMPLE.read_bytes())
        self.output = self.folder / "output" / "模拟运营报表.xlsx"

    def write(self, rows):
        with self.source.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=REQUIRED_FIELDS)
            writer.writeheader()
            writer.writerows(rows)

    def open_report(self):
        workbook = load_workbook(self.output, data_only=False)
        self.addCleanup(workbook.close)
        return workbook

    def assert_no_report_or_temporary(self):
        self.assertFalse(self.output.exists())
        self.assertEqual(list(self.folder.rglob("*.tmp")), [])

    def test_normal_export_creates_valid_three_sheet_xlsx(self):
        result = export_excel(self.source, self.output)
        self.assertEqual(result, self.output.absolute())
        self.assertEqual(self.open_report().sheetnames, list(SHEET_NAMES))

    def test_overview_exactly_matches_all_existing_metrics(self):
        report = calculate_metrics(self.source)
        export_excel(self.source, self.output)
        values = {r[0]: r[1] for r in self.open_report()["运营总览"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(values["商品明细数"], report.line_count)
        self.assertEqual(values["唯一订单数"], report.total_order_count)
        self.assertEqual(values["支付订单数"], report.paid_order_count)
        self.assertEqual(values["净有效订单数"], report.net_order_count)
        for label, cents in (("支付金额", report.payment_cents), ("退款金额", report.refund_cents),
                             ("退款后净销售额", report.net_sales_cents), ("支付客单价", report.paid_average_cents),
                             ("净客单价", report.net_average_cents)):
            self.assertEqual(Decimal(str(values[label])), Decimal(format_money(cents)))
        self.assertEqual(values["支付金额"], 3724.57)
        self.assertEqual(values["净客单价"], 100.85)

    def test_product_ranking_and_categories_reuse_metrics_order(self):
        report = calculate_metrics(self.source)
        export_excel(self.source, self.output)
        actual = list(self.open_report()["商品销量排行"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(actual, [(i, p.product_id, p.product_name, p.category, p.net_quantity)
                                  for i, p in enumerate(report.products, 1)])
        self.assertEqual(sum(r[4] for r in actual), 116)

    def test_summary_records_input_validation_counts_scope_and_time(self):
        before = datetime.now(BEIJING_TIME).replace(tzinfo=None, microsecond=0)
        export_excel(self.source, self.output)
        after = datetime.now(BEIJING_TIME).replace(tzinfo=None, microsecond=0)
        values = dict(self.open_report()["数据校验摘要"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(values["输入 CSV 文件名"], self.source.name)
        self.assertIn("整份数据通过", values["校验结果"])
        self.assertEqual((values["商品明细数"], values["唯一订单数"]), (80, 40))
        self.assertIn("净销售额", values["金额口径"])
        self.assertIn("不扣件数", values["商品排行口径"])
        self.assertIn("不是每日资金流水", values["数据时间边界"])
        # 不依赖电脑本地时区；工作簿明确标明 UTC+08:00，并保存日期值。
        self.assertLessEqual(before, values["报表生成时间（北京时间 UTC+08:00）"])
        self.assertLessEqual(values["报表生成时间（北京时间 UTC+08:00）"], after)

    def test_money_cells_are_numeric_with_two_decimal_format_and_yuan_units(self):
        export_excel(self.source, self.output)
        sheet = self.open_report()["运营总览"]
        for index in (4, 5, 6, 9, 10):
            self.assertEqual(sheet.cell(index, 2).data_type, "n")
            self.assertEqual(sheet.cell(index, 2).number_format, MONEY_FORMAT)
            self.assertTrue(sheet.cell(index, 3).value.startswith("元"))

    def test_headers_freeze_filter_widths_styles_and_chinese_text(self):
        export_excel(self.source, self.output)
        workbook = self.open_report()
        for sheet in workbook:
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertTrue(sheet["A1"].font.bold)
            self.assertEqual(sheet["A1"].fill.fgColor.rgb, "00244A64")
            self.assertGreaterEqual(sheet.column_dimensions["A"].width, 10)
        self.assertEqual(workbook["商品销量排行"].auto_filter.ref, "A1:E11")
        self.assertEqual(workbook["商品销量排行"]["C2"].value, "云朵桌面收纳盒")

    def test_no_formulas_macros_external_links_or_hyperlinks(self):
        export_excel(self.source, self.output)
        for sheet in self.open_report():
            for cells in sheet:
                for cell in cells:
                    self.assertNotEqual(cell.data_type, "f")
                    self.assertIsNone(cell.hyperlink)
        with ZipFile(self.output) as archive:
            self.assertFalse(any("vbaProject" in name or "externalLinks" in name for name in archive.namelist()))

    def test_equal_rank_ties_remain_product_id_string_order(self):
        self.write([row("A1", "P2"), row("B1", "P10")])
        export_excel(self.source, self.output)
        self.assertEqual([r[1] for r in self.open_report()["商品销量排行"].iter_rows(min_row=2, values_only=True)],
                         ["P10", "P2"])

    def test_half_cent_average_uses_existing_half_up_display_rounding(self):
        self.write([row(unit_price="1.00", payment_amount="1.00"),
                    row("B1", "P2", unit_price="1.01", payment_amount="1.01")])
        export_excel(self.source, self.output)
        sheet = self.open_report()["运营总览"]
        self.assertEqual(sheet["B4"].value, 2.01)
        self.assertEqual(sheet["B9"].value, 1.01)
        self.assertEqual(sheet["B10"].value, 1.01)

    def test_small_decimal_amounts_remain_exact_to_cents(self):
        self.write([row(unit_price="0.10", payment_amount="0.10"),
                    row("B1", "P2", unit_price="0.20", payment_amount="0.20")])
        export_excel(self.source, self.output)
        sheet = self.open_report()["运营总览"]
        self.assertEqual(Decimal(str(sheet["B4"].value)), Decimal("0.30"))
        self.assertEqual(Decimal(str(sheet["B9"].value)), Decimal("0.15"))

    def test_fifteen_digit_amount_including_cents_round_trips(self):
        self.write([row(unit_price="9999999999999.99", payment_amount="9999999999999.99")])
        export_excel(self.source, self.output)
        self.assertEqual(Decimal(str(self.open_report()["运营总览"]["B4"].value)), Decimal("9999999999999.99"))

    def test_larger_amount_is_rejected_without_changing_metrics_precision(self):
        amount = "123456789012345678901234567890.12"
        self.write([row(unit_price=amount, payment_amount=amount)])
        self.assertEqual(calculate_metrics(self.source).payment_cents, 12345678901234567890123456789012)
        with self.assertRaisesRegex(ExcelReportError, "金额.*安全数值精度"):
            export_excel(self.source, self.output)
        self.assert_no_report_or_temporary()

    def test_oversized_item_quantity_is_not_silently_rounded_by_excel(self):
        self.write([row(quantity="1000000000000000", unit_price="0.01",
                        discount_amount="9999999999999.99", payment_amount="0.01")])
        with self.assertRaisesRegex(ExcelReportError, "件数或订单数.*精度"):
            export_excel(self.source, self.output)
        self.assert_no_report_or_temporary()

    def test_zero_paid_orders_show_unavailable_averages_and_empty_ranking(self):
        self.write([row(order_status="未支付", payment_amount="0.00")])
        export_excel(self.source, self.output)
        workbook = self.open_report()
        self.assertEqual(workbook["运营总览"]["B4"].value, 0)
        self.assertIn("无法计算", workbook["运营总览"]["B9"].value)
        self.assertIn("无法计算", workbook["运营总览"]["B10"].value)
        self.assertEqual(workbook["商品销量排行"].max_row, 1)

    def test_full_refund_keeps_paid_order_but_net_average_is_unavailable(self):
        self.write([row(order_status="已全额退款", refund_amount="10.00", refund_quantity="1")])
        export_excel(self.source, self.output)
        sheet = self.open_report()["运营总览"]
        self.assertEqual([sheet.cell(i, 2).value for i in (4, 5, 6, 7, 8, 9)], [10, 10, 0, 1, 0, 10])
        self.assertIn("无法计算", sheet["B10"].value)

    def test_late_invalid_row_refuses_whole_export_and_does_not_create_output_dir(self):
        self.write([row(), row("B1", quantity="0")])
        with self.assertRaisesRegex(CsvValidationError, "第 3 行.*quantity"):
            export_excel(self.source, self.output)
        self.assertFalse(self.output.parent.exists())

    def test_missing_input_is_not_an_empty_success_report(self):
        with self.assertRaises(CsvValidationError):
            export_excel(self.folder / "不存在.csv", self.output)
        self.assertFalse(self.output.parent.exists())

    def test_missing_output_directory_is_created_only_after_validation(self):
        self.output = self.folder / "新目录" / "多层" / "新报表.xlsx"
        export_excel(self.source, self.output)
        self.assertEqual(self.open_report().sheetnames, list(SHEET_NAMES))

    def test_existing_file_is_not_overwritten_or_deleted(self):
        self.output.parent.mkdir()
        original = b"already-existing-file"
        self.output.write_bytes(original)
        with self.assertRaisesRegex(ExcelReportError, "拒绝覆盖"):
            export_excel(self.source, self.output)
        self.assertEqual(self.output.read_bytes(), original)
        self.assertEqual(list(self.folder.rglob("*.tmp")), [])

    def test_existing_directory_with_xlsx_name_is_not_modified(self):
        self.output.mkdir(parents=True)
        sentinel = self.output / "保留.txt"
        sentinel.write_text("保留", encoding="utf-8")
        with self.assertRaisesRegex(ExcelReportError, "拒绝覆盖"):
            export_excel(self.source, self.output)
        self.assertEqual(sentinel.read_text(encoding="utf-8"), "保留")

    def test_target_created_after_initial_check_is_not_overwritten(self):
        publish = excel_report._publish_no_replace

        def competing_writer(temporary, destination):
            destination.write_bytes(b"another-program-created-this")
            publish(temporary, destination)

        with patch("excel_report._publish_no_replace", side_effect=competing_writer):
            with self.assertRaises(ExcelReportError):
                export_excel(self.source, self.output)
        self.assertEqual(self.output.read_bytes(), b"another-program-created-this")
        self.assertEqual(list(self.folder.rglob("*.tmp")), [])

    def test_wrong_output_extension_is_rejected(self):
        with self.assertRaisesRegex(ExcelReportError, "扩展名"):
            export_excel(self.source, self.folder / "误用.csv")
        self.assertEqual(sorted(p.name for p in self.folder.iterdir()), [self.source.name])

    def test_parent_path_that_is_a_file_does_not_get_replaced(self):
        self.output.parent.write_bytes(b"not-a-directory")
        with self.assertRaises(ExcelReportError):
            export_excel(self.source, self.output)
        self.assertEqual(self.output.parent.read_bytes(), b"not-a-directory")

    def test_formula_like_names_and_error_codes_are_literal_text(self):
        for text in ('=HYPERLINK("https://example.invalid","点击")', "+1+1", "-1+1", "@SUM(1,2)", "#N/A"):
            with self.subTest(text=text):
                self.write([row(product_name=text)])
                target = self.folder / ("安全文本_" + str(len(list(self.folder.glob('安全文本_*.xlsx')))) + ".xlsx")
                export_excel(self.source, target)
                workbook = load_workbook(target, data_only=False)
                try:
                    cell = workbook["商品销量排行"]["C2"]
                    self.assertEqual((cell.value, cell.data_type, cell.hyperlink), (text, "s", None))
                finally:
                    workbook.close()

    def test_formula_like_product_id_and_input_filename_remain_text(self):
        self.write([row(product_id="=1+1")])
        self.source = self.source.rename(self.folder / "=SUM(1,2).csv")
        export_excel(self.source, self.output)
        workbook = self.open_report()
        for cell in (workbook["商品销量排行"]["B2"], workbook["数据校验摘要"]["B2"]):
            self.assertTrue(cell.value.startswith("="))
            self.assertEqual(cell.data_type, "s")

    def test_long_text_is_rejected_instead_of_silently_truncated(self):
        self.write([row(product_name="学" * 32768)])
        with self.assertRaisesRegex(ExcelReportError, "长度限制"):
            export_excel(self.source, self.output)
        self.assert_no_report_or_temporary()

    def test_xml_invalid_text_does_not_leave_a_formal_report(self):
        self.write([row(product_name="模拟\uffff商品")])
        with self.assertRaises(ExcelReportError):
            export_excel(self.source, self.output)
        self.assert_no_report_or_temporary()

    def test_serialization_exception_creates_no_output(self):
        with patch.object(Workbook, "save", side_effect=RuntimeError("模拟序列化失败")):
            with self.assertRaisesRegex(ExcelReportError, "模拟序列化失败"):
                export_excel(self.source, self.output)
        self.assert_no_report_or_temporary()

    def test_corrupt_serialization_is_rejected_by_readback_before_publishing(self):
        with patch.object(Workbook, "save", side_effect=lambda stream: stream.write(b"broken-xlsx")):
            with self.assertRaises(ExcelReportError):
                export_excel(self.source, self.output)
        self.assertFalse(self.output.parent.exists())

    def test_readback_detects_changed_money_value(self):
        original_save = Workbook.save

        def changed_save(workbook, stream):
            original_save(workbook, stream)
            stream.seek(0)
            changed = load_workbook(stream)
            try:
                changed["运营总览"]["B4"] = 0
                stream.seek(0)
                stream.truncate()
                original_save(changed, stream)
            finally:
                changed.close()

        with patch.object(Workbook, "save", new=changed_save):
            with self.assertRaisesRegex(ExcelReportError, "回读不一致"):
                export_excel(self.source, self.output)
        self.assert_no_report_or_temporary()

    def test_partial_disk_write_failure_cleans_owned_temporary_file(self):
        def partial_write(stream, payload):
            stream.write(payload[:20])
            raise OSError("模拟磁盘空间不足")

        with patch("excel_report._write_payload", side_effect=partial_write):
            with self.assertRaisesRegex(ExcelReportError, "磁盘空间不足"):
                export_excel(self.source, self.output)
        self.assert_no_report_or_temporary()

    def test_flush_failure_leaves_no_formal_report_or_temporary(self):
        with patch("excel_report.os.fsync", side_effect=OSError("模拟刷盘失败")):
            with self.assertRaisesRegex(ExcelReportError, "刷盘失败"):
                export_excel(self.source, self.output)
        self.assert_no_report_or_temporary()

    def test_publish_failure_preserves_inputs_and_removes_temporary(self):
        original = self.source.read_bytes()
        with patch("excel_report.os.rename", side_effect=PermissionError("模拟目标被占用")):
            with self.assertRaisesRegex(ExcelReportError, "目标被占用"):
                export_excel(self.source, self.output)
        self.assertEqual(self.source.read_bytes(), original)
        self.assert_no_report_or_temporary()

    def test_keyboard_interrupt_is_not_swallowed_and_temporary_is_cleaned(self):
        with patch("excel_report._write_payload", side_effect=KeyboardInterrupt):
            with self.assertRaises(KeyboardInterrupt):
                export_excel(self.source, self.output)
        self.assert_no_report_or_temporary()

    def test_cleanup_failure_reports_temporary_path_instead_of_success(self):
        with patch("excel_report._write_payload", side_effect=OSError("模拟写入失败")), \
                patch.object(Path, "unlink", side_effect=PermissionError("模拟清理受阻")):
            with self.assertRaisesRegex(ExcelReportError, "临时文件无法清理") as caught:
                export_excel(self.source, self.output)
        self.assertIn(".tmp", str(caught.exception))
        self.assertFalse(self.output.exists())
        self.assertEqual(len(list(self.folder.rglob("*.tmp"))), 1)

    def test_default_output_location_is_separate_and_names_are_unique(self):
        with patch("excel_report.DEFAULT_OUTPUT_DIR", self.folder / "默认输出"):
            first = export_excel(self.source)
            second = export_excel(self.source)
        self.assertEqual(first.parent, self.folder / "默认输出")
        self.assertNotEqual(first.name, second.name)
        self.assertTrue(first.exists() and second.exists())
        self.assertEqual(len(list(first.parent.iterdir())), 2)

    def test_export_uses_existing_metrics_once_and_does_not_modify_csv(self):
        original = self.source.read_bytes()
        sample_hash = hashlib.sha256(SAMPLE.read_bytes()).digest()
        with patch("excel_report.calculate_metrics", wraps=calculate_metrics) as calculate:
            export_excel(self.source, self.output)
        calculate.assert_called_once_with(self.source)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual(hashlib.sha256(SAMPLE.read_bytes()).digest(), sample_hash)

    def test_cli_success_and_validation_failure_messages(self):
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            self.assertEqual(main([str(self.source), "--output", str(self.output)]), 0)
        self.assertIn("Excel 报表已生成", output.getvalue())
        self.write([row(quantity="0")])
        failed = self.folder / "不应生成.xlsx"
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            self.assertEqual(main([str(self.source), "--output", str(failed)]), 1)
        self.assertIn("第 2 行 [quantity]", output.getvalue())
        self.assertIn("未生成 Excel 报表", output.getvalue())
        self.assertNotIn("Excel 报表已生成", output.getvalue())
        self.assertFalse(failed.exists())

    def test_cli_export_failure_never_prints_success(self):
        with patch("excel_report._write_payload", side_effect=OSError("模拟失败")):
            output = io.StringIO()
            with contextlib.redirect_stdout(output):
                self.assertEqual(main([str(self.source), "--output", str(self.output)]), 1)
        self.assertIn("Excel 导出未完成", output.getvalue())
        self.assertNotIn("Excel 报表已生成", output.getvalue())
        self.assert_no_report_or_temporary()

    def test_cli_default_sample_works_from_other_working_directory(self):
        result = subprocess.run([sys.executable, "-B", "-X", "utf8", str(ROOT / "report_excel.py"),
                                 "--output", str(self.output)], cwd=self.folder,
                                capture_output=True, text=True, encoding="utf-8", timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Excel 报表已生成", result.stdout)
        self.assertEqual(self.open_report()["运营总览"]["B4"].value, 3724.57)

    def test_import_does_not_read_csv_write_files_or_export(self):
        script = '''import openpyxl, csv_reader, order_validator, metrics
from unittest.mock import patch
with patch("pathlib.Path.read_bytes", side_effect=AssertionError("unexpected data read")), patch("pathlib.Path.write_bytes", side_effect=AssertionError("unexpected write")), patch("pathlib.Path.mkdir", side_effect=AssertionError("unexpected directory")), patch("openpyxl.Workbook.save", side_effect=AssertionError("unexpected export")):
    import excel_report, report_excel
print("IMPORT_SAFE")
'''
        result = subprocess.run([sys.executable, "-B", "-X", "utf8", "-c", script], cwd=ROOT,
                                capture_output=True, text=True, encoding="utf-8", timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("IMPORT_SAFE", result.stdout)


if __name__ == "__main__":
    unittest.main(verbosity=2)

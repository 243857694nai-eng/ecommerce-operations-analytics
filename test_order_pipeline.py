"""5B 端到端集成：所有输入副本及输出只写入自动清理的系统临时目录。"""

import contextlib
import hashlib
import io
import subprocess
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook

import excel_reader
import excel_report
import order_validator
from csv_reader import REQUIRED_FIELDS, read_csv
from excel_report import ExcelReportError, SHEET_NAMES, export_excel
from field_mapping import ALIASES
from metrics import calculate_metrics, format_money
from order_loader import InputCancelled, InputValidationError, load_orders
from report_excel import main as excel_main
from report_metrics import main as metrics_main


ROOT = Path(__file__).parent
SAMPLE = ROOT / "data" / "sample_orders.csv"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class OrderPipelineTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="order_pipeline_test_")
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)
        self.csv = self.folder / "模拟订单.csv"
        self.csv.write_bytes(SAMPLE.read_bytes())
        self.xlsx = self.folder / "模拟订单.xlsx"
        self.output = self.folder / "输出" / "运营报表.xlsx"
        self.rows = [row.values for row in read_csv(self.csv)]
        self.write_excel()

    def write_excel(self, rows=None, fields=REQUIRED_FIELDS, chinese=False, sheet_name="订单明细", header_row=3,
                    ambiguous=False, iso=False):
        workbook = Workbook()
        workbook.iso_dates = iso
        workbook.active.title = "说明页"
        workbook.active.append(["独立的虚构模拟订单，不作为订单数据"])
        sheet = workbook.create_sheet(sheet_name)
        for _ in range(header_row - 1):
            sheet.append(["说明：请明确选择表头行"])
        headers = [ALIASES[field][0] if chinese else field for field in fields]
        if ambiguous:
            headers[fields.index("payment_amount")] = "实付金额"
        sheet.append(headers)
        for values in (self.rows if rows is None else rows):
            sheet.append([values[field] for field in fields])
        workbook.save(self.xlsx)
        workbook.close()

    def answers(self, final="确认", count=14, header_row=3):
        return ["2", str(header_row)] + [""] * count + [final]

    def export(self, answers=None):
        with patch("builtins.input", side_effect=self.answers() if answers is None else answers), \
                contextlib.redirect_stdout(io.StringIO()):
            return export_excel(self.xlsx, self.output)

    def cli(self, entry=excel_main, answers=None, source=None):
        arguments = [str(self.xlsx if source is None else source)]
        if entry is excel_main:
            arguments += ["--output", str(self.output)]
        text = io.StringIO()
        with patch("builtins.input", side_effect=self.answers() if answers is None else answers), \
                contextlib.redirect_stdout(text):
            result = entry(arguments)
        return result, text.getvalue()

    def report_book(self):
        workbook = load_workbook(self.output, data_only=False)
        self.addCleanup(workbook.close)
        return workbook

    def assert_no_output(self):
        self.assertFalse(self.output.exists())
        self.assertFalse(self.output.parent.exists())
        self.assertEqual(list(self.folder.rglob("*.tmp")), [])

    def test_csv_default_and_existing_path_never_prompt_for_mapping(self):
        with patch("builtins.input", side_effect=AssertionError("CSV不应交互")):
            orders = load_orders(self.csv)
            report = calculate_metrics(self.csv)
            destination = export_excel(self.csv, self.output)
        self.assertEqual((len(orders.records), orders.source.file_format), (80, "CSV"))
        self.assertEqual(report.payment_cents, 372457)
        self.assertEqual(destination, self.output.absolute())
        summary = dict(self.report_book()["数据校验摘要"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary["输入 CSV 文件名"], self.csv.name)
        self.assertEqual(summary["输入文件格式"], "CSV")
        self.assertNotIn("输入工作表", summary)

    def test_csv_reader_does_not_gain_a_new_suffix_requirement(self):
        path = self.folder / "原有无后缀用法"
        path.write_bytes(self.csv.read_bytes())
        with patch("builtins.input", side_effect=AssertionError("不应交互")):
            self.assertEqual(calculate_metrics(path), calculate_metrics(self.csv))

    def test_excel_loader_returns_only_complete_validated_records_and_provenance(self):
        with patch("builtins.input", side_effect=self.answers()), contextlib.redirect_stdout(io.StringIO()):
            orders = load_orders(self.xlsx)
        self.assertIsInstance(orders.records, tuple)
        self.assertEqual(len(orders.records), 80)
        self.assertEqual((orders.source.file_name, orders.source.file_format, orders.source.sheet_name,
                          orders.source.header_row, orders.source.sha256),
                         (self.xlsx.name, "XLSX", "订单明细", 3, digest(self.xlsx)))

    def test_same_csv_and_excel_produce_identical_all_metrics_and_products(self):
        csv_report = calculate_metrics(self.csv)
        with patch("builtins.input", side_effect=self.answers()), contextlib.redirect_stdout(io.StringIO()):
            excel = calculate_metrics(self.xlsx)
        self.assertEqual(excel, csv_report)
        self.assertEqual(excel.products, csv_report.products)
        self.assertEqual((excel.payment_cents, excel.refund_cents, excel.net_sales_cents,
                          excel.paid_order_count, excel.net_order_count, excel.net_quantity),
                         (372457, 90086, 282371, 32, 28, 116))
        self.assertEqual((format_money(excel.paid_average_cents), format_money(excel.net_average_cents)),
                         ("116.39", "100.85"))
        self.assertNotEqual(excel.source, csv_report.source)

    def test_excel_export_contains_three_sheets_and_all_overview_values(self):
        self.export()
        workbook = self.report_book()
        self.assertEqual(workbook.sheetnames, list(SHEET_NAMES))
        actual = dict((row[0], row[1]) for row in workbook["运营总览"].iter_rows(min_row=2, values_only=True))
        reference = calculate_metrics(self.csv)
        self.assertEqual((actual["商品明细数"], actual["唯一订单数"], actual["支付订单数"], actual["净有效订单数"]),
                         (80, 40, 32, 28))
        for label, value in (("支付金额", reference.payment_cents), ("退款金额", reference.refund_cents),
                             ("退款后净销售额", reference.net_sales_cents), ("支付客单价", reference.paid_average_cents),
                             ("净客单价", reference.net_average_cents)):
            self.assertEqual(Decimal(str(actual[label])), Decimal(format_money(value)))

    def test_excel_export_keeps_product_sorting_and_styles(self):
        self.export()
        workbook = self.report_book()
        reference = calculate_metrics(self.csv)
        actual = list(workbook["商品销量排行"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(actual, [(i, row.product_id, row.product_name, row.category, row.net_quantity)
                                  for i, row in enumerate(reference.products, 1)])
        self.assertEqual([row[1] for row in actual[:3]], ["P003", "P001", "P007"])
        self.assertEqual(workbook["商品销量排行"].auto_filter.ref, "A1:E11")
        for sheet in workbook:
            self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(workbook["运营总览"]["B4"].number_format, "#,##0.00")

    def test_excel_summary_records_actual_sheet_header_format_and_snapshot(self):
        self.export()
        summary = dict(self.report_book()["数据校验摘要"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary["输入 Excel 文件名"], self.xlsx.name)
        self.assertEqual(summary["输入文件格式"], "XLSX")
        self.assertEqual(summary["输入工作表"], "订单明细")
        self.assertEqual(summary["表头原始行号"], 3)
        self.assertEqual(summary["输入快照 SHA256"], digest(self.xlsx))
        self.assertIn("内存快照", summary["快照说明"])
        self.assertNotIn("输入 CSV 文件名", summary)

    def test_reordered_chinese_headers_reuse_existing_mapping(self):
        self.write_excel(fields=tuple(reversed(REQUIRED_FIELDS)), chinese=True)
        self.export()
        self.assertEqual(self.report_book()["运营总览"]["B4"].value, 3724.57)

    def test_uppercase_xlsx_suffix_is_recognized(self):
        self.xlsx = self.xlsx.rename(self.folder / "模拟订单.XLSX")
        self.export()
        self.assertEqual(self.report_book()["运营总览"]["B6"].value, 2823.71)

    def test_mapping_cancellation_refuses_metrics_and_export(self):
        for entry in (metrics_main, excel_main):
            with self.subTest(entry=entry):
                code, text = self.cli(entry, self.answers(final="取消"))
                self.assertEqual(code, 2)
                self.assertIn("最终映射预览", text)
                self.assertNotIn("Excel 报表已生成", text)
                self.assertNotIn("运营指标报告", text)
        self.assert_no_output()

    def test_cancelled_mapping_does_not_call_validator_or_workbook_builder(self):
        with patch("excel_reader.validate_rows") as validator, patch("excel_report._build_workbook") as builder:
            with self.assertRaises(InputCancelled):
                self.export(self.answers(final="取消"))
        validator.assert_not_called()
        builder.assert_not_called()
        self.assert_no_output()

    def test_cancel_eof_and_keyboard_interrupt_all_stop_before_output(self):
        for answers in (["q"], [EOFError()], [KeyboardInterrupt()]):
            with self.subTest(answers=answers):
                code, text = self.cli(answers=answers)
                self.assertEqual(code, 2)
                self.assertNotIn("Excel 报表已生成", text)
                self.assert_no_output()

    def test_missing_required_field_never_enters_calculation_or_export(self):
        self.write_excel(fields=tuple(field for field in REQUIRED_FIELDS if field != "refund_quantity"))
        with patch("excel_report._build_workbook") as builder:
            code, text = self.cli(answers=self.answers(count=13))
        self.assertEqual(code, 1)
        self.assertIn("refund_quantity", text)
        builder.assert_not_called()
        self.assert_no_output()

    def test_ambiguous_amount_requires_extra_confirmation(self):
        self.write_excel(ambiguous=True)
        code, _ = self.cli(answers=["2", "3"] + [""] * 10 + ["不确定"])
        self.assertEqual(code, 2)
        self.assert_no_output()
        code, text = self.cli(answers=["2", "3"] + [""] * 10 + ["明细"] + [""] * 4 + ["确认"])
        self.assertEqual(code, 0, text)

    def test_last_row_invalid_amount_refuses_entire_input_not_partial_report(self):
        rows = [dict(row) for row in self.rows]
        rows[-1]["unit_price"] = "19.905"
        self.write_excel(rows)
        for entry in (metrics_main, excel_main):
            with self.subTest(entry=entry):
                code, text = self.cli(entry)
                self.assertEqual(code, 1)
                self.assertIn("第 83 行 H 列", text)
                self.assertNotIn("运营指标报告", text)
                self.assertNotIn("Excel 报表已生成", text)
        self.assert_no_output()

    def test_formula_rejected_before_metrics_or_export(self):
        rows = [dict(row) for row in self.rows]
        rows[-1]["payment_amount"] = "=1+1"
        self.write_excel(rows)
        code, text = self.cli()
        self.assertEqual(code, 1)
        self.assertIn("第 83 行 J 列", text)
        self.assertIn("公式", text)
        self.assert_no_output()

    def test_native_iso_timezone_and_invalid_suffix_never_reach_report(self):
        rows = [dict(self.rows[0])]
        rows[0]["order_time"] = datetime(2026, 10, 1, 10, 7, 11)
        for suffix in ("Z", "+09:00", "INVALID"):
            with self.subTest(suffix=suffix):
                self.write_excel(rows, iso=True)
                with zipfile.ZipFile(io.BytesIO(self.xlsx.read_bytes())) as archive:
                    parts = {name: archive.read(name) for name in archive.namelist()}
                namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                root = ElementTree.fromstring(parts["xl/worksheets/sheet2.xml"])
                root.find(f".//{namespace}c[@r='C4']/{namespace}v").text += suffix
                parts["xl/worksheets/sheet2.xml"] = ElementTree.tostring(root, encoding="utf-8")
                with zipfile.ZipFile(self.xlsx, "w") as archive:
                    for name, payload in parts.items():
                        archive.writestr(name, payload)
                before = digest(self.xlsx)
                code, text = self.cli(answers=[])
                self.assertEqual(code, 1)
                self.assertIn("第 4 行 C 列", text)
                self.assertEqual(digest(self.xlsx), before)
                self.assert_no_output()

    def test_csv_and_excel_each_use_their_existing_validator_once(self):
        with patch("order_validator.validate_rows", wraps=order_validator.validate_rows) as csv_validator:
            calculate_metrics(self.csv)
        csv_validator.assert_called_once()
        with patch("excel_reader.validate_rows", wraps=excel_reader.validate_rows) as excel_validator:
            self.export()
        excel_validator.assert_called_once()

    def test_export_calls_existing_metrics_once_and_reads_excel_once(self):
        with patch("excel_report.calculate_metrics", wraps=calculate_metrics) as calculate, \
                patch("excel_reader.read_snapshot", wraps=excel_reader.read_snapshot) as reader:
            self.export()
        calculate.assert_called_once_with(self.xlsx)
        reader.assert_called_once_with(self.xlsx)

    def test_same_name_output_is_not_overwritten(self):
        self.output.parent.mkdir()
        self.output.write_bytes(b"preserve-existing-report")
        before = digest(self.output)
        with self.assertRaisesRegex(ExcelReportError, "拒绝覆盖"):
            self.export()
        self.assertEqual(digest(self.output), before)
        self.assertEqual(list(self.folder.rglob("*.tmp")), [])

    def test_input_excel_cannot_be_used_as_output_destination(self):
        self.output = self.xlsx
        before = digest(self.xlsx)
        with self.assertRaisesRegex(ExcelReportError, "拒绝覆盖"):
            self.export()
        self.assertEqual(digest(self.xlsx), before)

    def test_target_created_during_publish_is_preserved(self):
        publish = excel_report._publish_no_replace
        def competing_writer(temporary, destination):
            destination.write_bytes(b"another-writer")
            publish(temporary, destination)
        with patch("excel_report._publish_no_replace", side_effect=competing_writer), self.assertRaises(ExcelReportError):
            self.export()
        self.assertEqual(self.output.read_bytes(), b"another-writer")
        self.assertEqual(list(self.folder.rglob("*.tmp")), [])

    def test_export_failure_is_not_reported_as_success_and_cleans_temp(self):
        with patch("excel_report._write_payload", side_effect=OSError("模拟写入失败")):
            code, text = self.cli()
        self.assertEqual(code, 1)
        self.assertIn("Excel 导出未完成", text)
        self.assertNotIn("Excel 报表已生成", text)
        self.assertFalse(self.output.exists())
        self.assertEqual(list(self.folder.rglob("*.tmp")), [])

    def test_serialization_failure_does_not_create_output_directory(self):
        with patch("excel_report._serialize_and_verify", side_effect=RuntimeError("模拟序列化失败")):
            code, text = self.cli()
        self.assertEqual(code, 1)
        self.assertNotIn("Excel 报表已生成", text)
        self.assert_no_output()

    def test_unexpected_validator_error_stops_both_cli_without_success(self):
        for entry in (metrics_main, excel_main):
            with self.subTest(entry=entry), patch("excel_reader.validate_rows", side_effect=RuntimeError("模拟内部异常")):
                code, text = self.cli(entry)
                self.assertEqual(code, 1)
                self.assertIn("RuntimeError", text)
                self.assertNotIn("运营指标报告", text)
                self.assertNotIn("Excel 报表已生成", text)
        self.assert_no_output()

    def test_api_preserves_error_chain_and_never_returns_invalid_data(self):
        rows = [dict(self.rows[0], order_status="未知状态")]
        self.write_excel(rows)
        with self.assertRaises(InputValidationError) as caught:
            self.export()
        self.assertIsNotNone(caught.exception.__cause__)
        self.assertFalse(hasattr(caught.exception, "records"))
        self.assert_no_output()

    def test_source_hashes_and_existing_independent_report_stay_unchanged(self):
        sentinel = self.folder / "先前报表.xlsx"
        sentinel.write_bytes(b"independent-existing-report")
        before = {path: digest(path) for path in (SAMPLE, self.csv, self.xlsx, sentinel)}
        self.export()
        self.assertEqual({path: digest(path) for path in before}, before)
        with self.assertRaises(InputCancelled):
            self.export(self.answers(final="取消"))
        self.assertEqual({path: digest(path) for path in before}, before)

    def test_default_export_still_uses_existing_output_directory_policy(self):
        with patch("excel_report.DEFAULT_OUTPUT_DIR", self.folder / "默认输出"), \
                patch("builtins.input", side_effect=self.answers()), contextlib.redirect_stdout(io.StringIO()):
            destination = export_excel(self.xlsx)
        self.assertEqual(destination.parent, self.folder / "默认输出")
        self.assertTrue(destination.name.startswith("运营报表_"))

    def test_source_filename_and_sheet_formula_like_text_remain_literal(self):
        self.write_excel(sheet_name="=1+1")
        self.xlsx = self.xlsx.rename(self.folder / "=SUM(1,2).xlsx")
        self.export()
        sheet = self.report_book()["数据校验摘要"]
        for coordinate in ("B2", "B13"):
            self.assertTrue(sheet[coordinate].value.startswith("="))
            self.assertEqual(sheet[coordinate].data_type, "s")
            self.assertIsNone(sheet[coordinate].hyperlink)

    def test_source_provenance_belongs_to_original_snapshot_after_file_change(self):
        before = digest(self.xlsx)
        answers = iter(self.answers())
        def answer(prompt):
            result = next(answers)
            if "开始整份校验" in prompt:
                self.xlsx.write_bytes(b"newer-file-version-in-test-only")
            return result
        with patch("builtins.input", side_effect=answer), contextlib.redirect_stdout(io.StringIO()):
            export_excel(self.xlsx, self.output)
        summary = dict(self.report_book()["数据校验摘要"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary["输入快照 SHA256"], before)
        self.assertEqual(self.xlsx.read_bytes(), b"newer-file-version-in-test-only")
        self.assertEqual(self.report_book()["运营总览"]["B4"].value, 3724.57)

    def test_metrics_cli_excel_prints_matching_numbers_and_source_only_after_validation(self):
        code, text = self.cli(metrics_main)
        self.assertEqual(code, 0)
        for expected in ("支付金额：3724.57 元", "净销售额：2823.71 元", "净客单价：100.85 元/笔",
                         "商品净成交件数合计：116 件", "输入工作表：'订单明细'", digest(self.xlsx)):
            self.assertIn(expected, text)
        self.assertNotIn("计算只读入 CSV", text)
        self.assert_no_output()

    def test_render_failure_is_not_reported_as_success(self):
        with patch("report_metrics.render_report", side_effect=RuntimeError("模拟渲染失败")):
            code, text = self.cli(metrics_main)
        self.assertEqual(code, 1)
        self.assertNotIn("运营指标报告", text)
        self.assert_no_output()

    def test_csv_pipeline_works_without_importing_excel_dependencies(self):
        script = '''import builtins
original = builtins.__import__
def guarded(name, *args, **kwargs):
    if name.startswith(('openpyxl', 'excel_reader', 'inspect_excel')):
        raise AssertionError('CSV不应导入Excel组件')
    return original(name, *args, **kwargs)
builtins.__import__ = guarded
from metrics import calculate_metrics
import sys
assert calculate_metrics(sys.argv[1]).payment_cents == 372457
print('CSV_STANDARD_LIBRARY_OK')
'''
        result = subprocess.run([sys.executable, "-B", "-X", "utf8", "-c", script, str(self.csv)], cwd=ROOT,
                                capture_output=True, text=True, encoding="utf-8", timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_new_modules_import_without_io_prompts_or_exports(self):
        script = '''from unittest.mock import patch
with patch('pathlib.Path.read_bytes', side_effect=AssertionError('read')), \\
     patch('pathlib.Path.write_bytes', side_effect=AssertionError('write')), \\
     patch('builtins.input', side_effect=AssertionError('prompt')):
    import order_loader, inspect_excel, metrics, report_metrics, excel_report, report_excel
print('SAFE_IMPORTS')
'''
        result = subprocess.run([sys.executable, "-B", "-X", "utf8", "-c", script], cwd=ROOT,
                                capture_output=True, text=True, encoding="utf-8", timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_real_excel_report_cli_process_with_simulated_keyboard_input(self):
        result = subprocess.run([sys.executable, "-B", "-X", "utf8", str(ROOT / "report_excel.py"), str(self.xlsx),
                                 "--output", str(self.output)], input="\n".join(self.answers()) + "\n", cwd=self.folder,
                                capture_output=True, text=True, encoding="utf-8", timeout=30)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("Excel 报表已生成", result.stdout)
        self.assertEqual(self.report_book().sheetnames, list(SHEET_NAMES))

    def test_real_cancelled_cli_process_has_nonzero_status_and_no_report(self):
        result = subprocess.run([sys.executable, "-B", "-X", "utf8", str(ROOT / "report_excel.py"), str(self.xlsx),
                                 "--output", str(self.output)], input="\n".join(self.answers(final="取消")) + "\n",
                                cwd=self.folder, capture_output=True, text=True, encoding="utf-8", timeout=30)
        self.assertEqual(result.returncode, 2, result.stdout + result.stderr)
        self.assertNotIn("Excel 报表已生成", result.stdout)
        self.assert_no_output()


if __name__ == "__main__":
    unittest.main()

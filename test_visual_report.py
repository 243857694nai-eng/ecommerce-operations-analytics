"""5C：隔离模拟文件、图表引用、口径和安全发布测试，不写原始样本。"""

import contextlib
import csv
import hashlib
import io
import subprocess
import sys
import tempfile
import unittest
import zipfile
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart

import excel_report
from csv_reader import CsvValidationError, REQUIRED_FIELDS
from excel_report import ExcelReportError, SHEET_NAMES, VISUAL_SHEET_NAMES, export_excel
from metrics import calculate_metrics
from report_excel import main
from test_excel_report import row


ROOT = Path(__file__).parent
SAMPLE = ROOT / "data" / "sample_orders.csv"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class VisualReportTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="visual_orders_test_")
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)
        self.source = self.folder / "模拟订单.csv"
        self.source.write_bytes(SAMPLE.read_bytes())
        self.output = self.folder / "output" / "模拟分析.xlsx"

    def write(self, rows):
        with self.source.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=REQUIRED_FIELDS)
            writer.writeheader()
            writer.writerows(rows)

    def xlsx(self):
        path = self.folder / "模拟订单.xlsx"
        workbook = Workbook()
        workbook.active.title = "说明页"
        workbook.active.append(["模拟数据，不选择此页"])
        sheet = workbook.create_sheet("模拟明细")
        sheet.append(["模拟订单，表头在第2行"])
        with self.source.open(encoding="utf-8-sig", newline="") as stream:
            for record in csv.reader(stream):
                sheet.append(record)
        workbook.save(path)
        workbook.close()
        return path

    def book(self):
        book = load_workbook(self.output, data_only=False)
        self.addCleanup(book.close)
        return book

    def export(self):
        return export_excel(self.source, self.output, visual=True)

    def answers(self, final="确认"):
        return ["2", "2"] + [""] * len(REQUIRED_FIELDS) + [final]

    def no_output(self):
        self.assertFalse(self.output.exists())
        self.assertEqual(list(self.folder.rglob("*.tmp")), [])

    def test_default_report_still_has_original_three_sheets(self):
        export_excel(self.source, self.output)
        self.assertEqual(self.book().sheetnames, list(SHEET_NAMES))

    def test_visual_report_has_six_sheets_and_three_native_charts(self):
        self.export()
        book = self.book()
        self.assertEqual(book.sheetnames, list(SHEET_NAMES + VISUAL_SHEET_NAMES))
        self.assertEqual([len(book[name]._charts) for name in VISUAL_SHEET_NAMES], [1, 1, 1])
        self.assertIsInstance(book["销售趋势"]._charts[0], LineChart)
        self.assertIsInstance(book["商品销售分析"]._charts[0], BarChart)

    def test_original_metrics_and_complete_product_ranking_are_preserved(self):
        self.export()
        book = self.book()
        overview = dict((r[0], r[1]) for r in book["运营总览"].iter_rows(min_row=2, values_only=True))
        self.assertEqual([overview[key] for key in ("商品明细数", "唯一订单数", "支付订单数", "净有效订单数")],
                         [80, 40, 32, 28])
        self.assertEqual([Decimal(str(overview[key])) for key in ("支付金额", "退款金额", "退款后净销售额",
                                                                  "支付客单价", "净客单价")],
                         list(map(Decimal, ("3724.57", "900.86", "2823.71", "116.39", "100.85"))))
        report = calculate_metrics(self.source)
        actual = list(book["商品销售分析"].iter_rows(min_row=5, values_only=True))
        self.assertEqual([(r[1], r[6], Decimal(str(r[9]))) for r in actual],
                         [(p.product_id, p.net_quantity, Decimal(p.net_sales_cents) / 100) for p in report.products])
        self.assertEqual(sum(r[4] for r in book["商品销量排行"].iter_rows(min_row=2, values_only=True)), 116)

    def test_daily_values_reconcile_to_independent_decimal_csv_grouping(self):
        expected = defaultdict(lambda: [Decimal(0), Decimal(0)])
        with self.source.open(encoding="utf-8-sig", newline="") as stream:
            for record in csv.DictReader(stream):
                amounts = expected[record["order_time"][:10]]
                if record["order_status"] in {"已支付", "部分退款", "已全额退款"}:
                    amounts[0] += Decimal(record["payment_amount"])
                    amounts[1] += Decimal(record["refund_amount"])
        self.export()
        actual = {}
        for day, paid, refund, net in self.book()["销售趋势"].iter_rows(min_row=5, values_only=True):
            actual[day.strftime("%Y-%m-%d")] = [Decimal(str(paid)), Decimal(str(refund))]
            self.assertEqual(Decimal(str(net)), Decimal(str(paid)) - Decimal(str(refund)))
        self.assertEqual(actual, dict(expected))
        self.assertEqual(sum(x[0] for x in actual.values()), Decimal("3724.57"))
        self.assertEqual(sum(x[1] for x in actual.values()), Decimal("900.86"))

    def test_sample_refund_counts_and_exact_ratios(self):
        report = calculate_metrics(self.source)
        self.assertEqual((report.refund_order_count, report.full_refund_order_count, report.partial_refund_order_count),
                         (8, 4, 4))
        self.assertEqual(report.refund_amount_ratio, Fraction(90086, 372457))
        self.assertEqual(report.refund_order_ratio, Fraction(1, 4))
        self.export()
        sheet = self.book()["退款分析"]
        self.assertEqual(sheet["B10"].value, 0.2419)
        self.assertEqual(sheet["B11"].value, 0.25)
        self.assertEqual(sheet["B10"].number_format, "0.00%")
        self.assertIn("支付金额", sheet["C10"].value)
        self.assertIn("支付订单数", sheet["C11"].value)

    def test_multiline_partial_refund_order_counts_once_and_price_refund_keeps_units(self):
        self.write([row(order_status="部分退款", refund_amount="1.00"),
                    row(product_id="P2", order_line_id="A1-02", order_status="部分退款", refund_amount="0.00")])
        report = calculate_metrics(self.source)
        self.assertEqual((report.paid_order_count, report.refund_order_count, report.partial_refund_order_count), (1, 1, 1))
        self.assertEqual((report.daily[0].paid_order_count, report.daily[0].refund_order_count), (1, 1))
        self.assertEqual(report.net_quantity, 2)
        self.assertEqual(report.refund_amount_ratio, Fraction(1, 20))

    def test_daily_sorting_and_missing_days_are_not_fabricated(self):
        self.write([row("B1", order_time="2026-10-03 10:00:00"), row(order_time="2026-10-01 10:00:00")])
        report = calculate_metrics(self.source)
        self.assertEqual([d.day.isoformat() for d in report.daily], ["2026-10-01", "2026-10-03"])
        self.assertEqual(sum(d.payment_cents for d in report.daily), report.payment_cents)

    def test_single_day_chart_has_one_real_point_with_marker(self):
        self.write([row()])
        self.export()
        chart = self.book()["销售趋势"]._charts[0]
        self.assertEqual(chart.series[0].val.numRef.f, "'销售趋势'!$B$5")
        self.assertEqual(chart.series[0].marker.symbol, "circle")
        self.assertEqual(chart.x_axis.numFmt.formatCode, "yyyy-mm-dd")

    def test_zero_refunds_have_zero_rates_and_zero_refund_chart_data(self):
        self.write([row()])
        self.export()
        sheet = self.book()["退款分析"]
        self.assertEqual([sheet[f"B{r}"].value for r in (5, 7, 8, 9, 10, 11)], [0, 0, 0, 0, 0, 0])
        self.assertEqual(sheet["B16"].value, 0)
        self.assertEqual(len(sheet._charts), 1)

    def test_unpaid_cancelled_only_have_no_product_chart_and_undefined_ratios(self):
        self.write([row(order_status="未支付", payment_amount="0.00"),
                    row("B1", order_status="已取消", payment_amount="0.00")])
        self.export()
        book = self.book()
        self.assertEqual(book["商品销售分析"]._charts, [])
        self.assertIn("没有已支付商品", book["商品销售分析"]["A5"].value)
        self.assertIn("无法计算", book["退款分析"]["B10"].value)
        self.assertIn("无法计算", book["退款分析"]["B11"].value)
        self.assertEqual(book["销售趋势"]["B5"].value, 0)

    def test_fully_refunded_orders_keep_payment_and_zero_net(self):
        self.write([row(order_status="已全额退款", refund_amount="10.00", refund_quantity="1")])
        self.export()
        book = self.book()
        self.assertEqual(book["销售趋势"]["D5"].value, 0)
        self.assertEqual(book["退款分析"]["B10"].value, 1)
        self.assertEqual(book["商品销售分析"]["G5"].value, 0)

    def test_empty_input_is_rejected_not_presented_as_zero_business(self):
        self.write([])
        with self.assertRaises(CsvValidationError):
            self.export()
        self.no_output()

    def test_late_illegal_row_rejects_entire_visual_report(self):
        self.write([row(), row("B1", quantity="0")])
        with self.assertRaises(CsvValidationError):
            self.export()
        self.no_output()

    def test_line_chart_series_reference_exact_dates_and_money_table(self):
        self.export()
        chart = self.book()["销售趋势"]._charts[0]
        self.assertEqual([s.val.numRef.f for s in chart.series],
                         [f"'销售趋势'!${c}$5:${c}$14" for c in "BCD"])
        self.assertEqual([s.tx.strRef.f for s in chart.series], [f"'销售趋势'!{c}4" for c in "BCD"])
        self.assertTrue(all(s.cat.numRef.f == "'销售趋势'!$A$5:$A$14" for s in chart.series))

    def test_refund_chart_reference_matches_daily_refund_table(self):
        self.export()
        chart = self.book()["退款分析"]._charts[0]
        self.assertEqual(chart.series[0].val.numRef.f, "'退款分析'!$B$16:$B$25")
        self.assertEqual(chart.series[0].cat.numRef.f, "'退款分析'!$A$16:$A$25")
        report = calculate_metrics(self.source)
        actual = list(self.book()["退款分析"].iter_rows(min_row=16, values_only=True))
        self.assertEqual([(r[2], r[3]) for r in actual],
                         [(d.refund_order_count, d.paid_order_count) for d in report.daily])

    def test_chart_axes_and_titles_are_explicitly_visible_without_legend_overlay(self):
        self.export()
        book = self.book()
        for name in VISUAL_SHEET_NAMES:
            chart = book[name]._charts[0]
            for axis in (chart.x_axis, chart.y_axis):
                self.assertFalse(axis.delete)
                self.assertEqual(axis.tickLblPos, "nextTo")
                self.assertIsNone(axis.title)
            self.assertFalse(chart.title.overlay)
        for name in ("销售趋势", "退款分析"):
            chart = book[name]._charts[0]
            self.assertFalse(chart.varyColors)
            self.assertFalse(chart.legend.overlay)
            self.assertEqual(chart.x_axis.axPos, "b")

    def test_top_ten_chart_preserves_full_table_over_ten_products(self):
        self.write([row(f"A{i}", f"P{i:02d}") for i in range(12)])
        self.export()
        sheet = self.book()["商品销售分析"]
        self.assertEqual(sheet.max_row, 16)
        self.assertEqual(sheet._charts[0].series[0].val.numRef.f, "'商品销售分析'!$G$5:$G$14")
        self.assertEqual(sheet._charts[0].series[0].cat.numRef.f, "'商品销售分析'!$K$5:$K$14")

    def test_same_name_different_ids_and_ties_remain_distinct(self):
        self.write([row("A1", "P2", product_name="同名模拟商品"), row("B1", "P10", product_name="同名模拟商品")])
        self.export()
        sheet = self.book()["商品销售分析"]
        self.assertEqual((sheet["B5"].value, sheet["B6"].value), ("P10", "P2"))
        self.assertNotEqual(sheet["K5"].value, sheet["K6"].value)

    def test_dates_money_headers_notes_and_chart_anchors(self):
        self.export()
        book = self.book()
        for name in VISUAL_SHEET_NAMES:
            sheet = book[name]
            self.assertIn("模拟数据", sheet["A1"].value)
            self.assertTrue(sheet["A4"].font.bold)
            self.assertEqual(sheet.freeze_panes, "A5")
            self.assertIn("虚构模拟", sheet["A2"].value)
        self.assertEqual(book["销售趋势"]["A5"].number_format, "yyyy-mm-dd")
        self.assertEqual(book["商品销售分析"]["J5"].number_format, "#,##0.00")
        self.assertEqual(book["销售趋势"]._charts[0].anchor._from.col, 5)
        self.assertEqual(book["商品销售分析"]._charts[0].anchor._from.col, 12)
        self.assertIn("不是每日实际退款流水", book["退款分析"]["A2"].value)

    def test_formula_like_product_names_are_text_not_formulas_or_external_links(self):
        self.write([row(product_name="=HYPERLINK(模拟)")])
        self.export()
        for sheet in self.book():
            for cells in sheet:
                for cell in cells:
                    self.assertNotEqual(cell.data_type, "f")
                    self.assertIsNone(cell.hyperlink)
        self.assertEqual(self.book()["商品销售分析"]["C5"].data_type, "s")
        with zipfile.ZipFile(self.output) as archive:
            self.assertFalse(any("externalLinks" in n for n in archive.namelist()))

    def test_excel_cli_generates_visual_report_with_accurate_provenance(self):
        xlsx = self.xlsx()
        before = {p: digest(p) for p in (SAMPLE, self.source, xlsx)}
        result = subprocess.run([sys.executable, "-B", "-X", "utf8", str(ROOT / "report_excel.py"), str(xlsx),
                                 "--visual", "--output", str(self.output)], cwd=self.folder,
                                input="\n".join(self.answers()) + "\n", capture_output=True,
                                text=True, encoding="utf-8", timeout=30)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("Excel 报表已生成", result.stdout)
        summary = dict(self.book()["数据校验摘要"].iter_rows(min_row=2, values_only=True))
        self.assertEqual((summary["输入工作表"], summary["表头原始行号"], summary["输入快照 SHA256"]),
                         ("模拟明细", 2, before[xlsx]))
        self.assertEqual({p: digest(p) for p in before}, before)
        with patch("builtins.input", side_effect=self.answers()), contextlib.redirect_stdout(io.StringIO()):
            excel = calculate_metrics(xlsx)
        self.assertEqual(excel, calculate_metrics(self.source))

    def test_excel_cancel_never_builds_visual_tables_or_reports(self):
        xlsx = self.xlsx()
        with patch("builtins.input", side_effect=self.answers("取消")), \
                patch("excel_report._add_visual_analysis") as builder, contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main([str(xlsx), "--visual", "--output", str(self.output)]), 2)
        builder.assert_not_called()
        self.no_output()

    def test_excel_formula_is_refused_without_partial_report(self):
        xlsx = self.xlsx()
        book = load_workbook(xlsx)
        book["模拟明细"]["J3"] = "=1+1"
        book.save(xlsx)
        book.close()
        before = digest(xlsx)
        with patch("builtins.input", side_effect=self.answers()), contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main([str(xlsx), "--visual", "--output", str(self.output)]), 1)
        self.assertEqual(digest(xlsx), before)
        self.no_output()

    def test_invalid_native_iso_date_is_rejected_before_confirmation_and_visual_export(self):
        xlsx = self.xlsx()
        with zipfile.ZipFile(xlsx) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        xml = ElementTree.fromstring(parts["xl/worksheets/sheet2.xml"])
        cell = xml.find(f".//{namespace}c[@r='C3']")
        cell.clear()
        cell.attrib.update(r="C3", t="d")
        ElementTree.SubElement(cell, namespace + "v").text = "2026-10-01T10:07:11Z"
        parts["xl/worksheets/sheet2.xml"] = ElementTree.tostring(xml)
        with zipfile.ZipFile(xlsx, "w") as archive:
            for name, payload in parts.items():
                archive.writestr(name, payload)
        before = digest(xlsx)
        with patch("builtins.input", side_effect=AssertionError("日期应先被拒绝")), contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main([str(xlsx), "--visual", "--output", str(self.output)]), 1)
        self.assertEqual(digest(xlsx), before)
        self.no_output()

    def test_existing_target_is_not_overwritten(self):
        self.output.parent.mkdir()
        self.output.write_bytes(b"existing-report")
        before = digest(self.output)
        with self.assertRaisesRegex(ExcelReportError, "拒绝覆盖"):
            self.export()
        self.assertEqual(digest(self.output), before)

    def test_target_created_during_publication_is_preserved(self):
        original = excel_report._publish_no_replace
        def competing_writer(temporary, destination):
            destination.write_bytes(b"concurrent-report")
            original(temporary, destination)
        with patch("excel_report._publish_no_replace", side_effect=competing_writer), self.assertRaises(ExcelReportError):
            self.export()
        self.assertEqual(self.output.read_bytes(), b"concurrent-report")
        self.assertEqual(list(self.folder.rglob("*.tmp")), [])

    def test_write_failure_is_not_success_and_owned_temporary_is_cleaned(self):
        text = io.StringIO()
        before = digest(self.source)
        with patch("excel_report._write_payload", side_effect=OSError("模拟写入失败")), contextlib.redirect_stdout(text):
            self.assertEqual(main([str(self.source), "--visual", "--output", str(self.output)]), 1)
        self.assertNotIn("Excel 报表已生成", text.getvalue())
        self.assertEqual(digest(self.source), before)
        self.no_output()

    def test_missing_chart_is_detected_before_any_output_file(self):
        save = Workbook.save
        def without_chart(book, stream):
            charts = book["销售趋势"]._charts
            book["销售趋势"]._charts = []
            try:
                save(book, stream)
            finally:
                book["销售趋势"]._charts = charts
        with patch.object(Workbook, "save", without_chart), self.assertRaisesRegex(ExcelReportError, "图表"):
            self.export()
        self.no_output()

    def test_changed_chart_reference_is_detected_before_publication(self):
        save = Workbook.save
        def changed_reference(book, stream):
            series = book["销售趋势"]._charts[0].series[0]
            original = series.val.numRef.f
            series.val.numRef.f = "'销售趋势'!$D$5:$D$14"
            try:
                save(book, stream)
            finally:
                series.val.numRef.f = original
        with patch.object(Workbook, "save", changed_reference), self.assertRaisesRegex(ExcelReportError, "图表"):
            self.export()
        self.no_output()

    def test_visual_build_failure_leaves_no_report_and_does_not_claim_success(self):
        text = io.StringIO()
        with patch("excel_report._add_visual_analysis", side_effect=RuntimeError("模拟图表构建失败")), \
                contextlib.redirect_stdout(text):
            self.assertEqual(main([str(self.source), "--visual", "--output", str(self.output)]), 1)
        self.assertNotIn("Excel 报表已生成", text.getvalue())
        self.no_output()

    def test_refund_ratios_round_only_for_display_to_one_hundredth_percent(self):
        self.write([row(order_status="已全额退款", refund_amount="10.00", refund_quantity="1"),
                    row("B1"), row("C1")])
        report = calculate_metrics(self.source)
        self.assertEqual(report.refund_amount_ratio, Fraction(1, 3))
        self.assertEqual(report.refund_order_ratio, Fraction(1, 3))
        self.export()
        self.assertEqual(self.book()["退款分析"]["B10"].value, 0.3333)
        self.assertEqual(self.book()["退款分析"]["B11"].value, 0.3333)

    def test_small_decimal_amounts_are_exact_and_source_is_unchanged(self):
        self.write([row(unit_price="0.10", payment_amount="0.10"),
                    row("B1", "P2", unit_price="0.20", payment_amount="0.20")])
        before = digest(self.source)
        self.export()
        self.assertEqual(Decimal(str(self.book()["销售趋势"]["B5"].value)), Decimal("0.30"))
        self.assertEqual(digest(self.source), before)

    def test_original_metrics_and_input_snapshot_are_computed_only_once(self):
        xlsx = self.xlsx()
        from excel_reader import read_snapshot
        with patch("excel_reader.read_snapshot", wraps=read_snapshot) as reader, \
                patch("excel_report.calculate_metrics", wraps=calculate_metrics) as calculator, \
                patch("builtins.input", side_effect=self.answers()), contextlib.redirect_stdout(io.StringIO()):
            export_excel(xlsx, self.output, visual=True)
        reader.assert_called_once_with(xlsx)
        calculator.assert_called_once_with(xlsx)

    def test_default_output_visual_mode_uses_existing_unique_filename_policy(self):
        with patch("excel_report.DEFAULT_OUTPUT_DIR", self.folder / "reports"):
            first = export_excel(self.source, visual=True)
            second = export_excel(self.source, visual=True)
        self.assertNotEqual(first, second)
        self.assertTrue(first.exists() and second.exists())


if __name__ == "__main__":
    unittest.main()
